"""Neutral locator inventory and separately bound original-target handoff.

Consumes the original paper and its reference layout, never enriched DOI/URL
fields. Callers retain upload authorization/safety responsibility. The inventory
remains text-free; the separate handoff returns original URLs to normal retrieval,
without requests, style verdicts or source-admission behavior here.
"""
import hashlib
import io
import json
import re
import zipfile
from typing import Literal
from urllib.parse import urlsplit

import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pydantic import BaseModel, ConfigDict, Field, model_validator

from app.services.ref_field_extractor import _DOI_PATTERN, _URL_PATTERN
from app.services.reference_layout import ReferenceLayoutArtifact, ReferenceLayoutRectangle
from app.services.schemas import ParsedReference
from app.services.schemas import SubmittedHyperlinkBinding

_LOCATOR_HINT = re.compile(
    r"\b(?:https?\s*:|https?\s*/|www\s*\.|doi\s*[:/]|doi\s*\.\s*org|10\.\d{4,}\s*/)"
    r"|(?<![\w@])(?:[a-z0-9-]+\.)+[a-z]{2,}/[^\s]*"
    r"|(?<![\w@])(?:[a-z0-9-]+\.)+(?:com|org|net|edu|gov)(?![\w])", re.I,
)


class SubmittedLocatorEntry(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    reference_id: str
    reference_text_sha256: str
    status: Literal["supplied", "not_observed", "unknown"]
    evidence_channels: tuple[Literal["raw_text", "visible_text", "native_hyperlink"], ...] = ()
    rectangles: tuple[ReferenceLayoutRectangle, ...] = ()
    paragraph_indexes: tuple[int, ...] = ()
    limitation: str | None = None
    locator_validity: Literal["not_assessed"] = "not_assessed"
    style_requirement: Literal["not_assessed"] = "not_assessed"
    contributes_to_issue_counts: Literal[False] = False


class SubmittedLocatorInventory(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    version: Literal["submitted-locator-inventory-v1", "submitted-locator-inventory-v2"] = "submitted-locator-inventory-v2"
    paper_sha256: str
    submitted_snapshot_sha256: str
    reference_list_coverage: Literal["complete", "partial", "unknown"] = "unknown"
    entries: tuple[SubmittedLocatorEntry, ...]
    total_entries: int = Field(ge=0)
    assessed_entries: int = Field(ge=0)
    counts: dict[str, int]
    contributes_to_issue_counts: Literal[False] = False

    @model_validator(mode="after")
    def validate_counts(self):
        expected = {k: sum(e.status == k for e in self.entries)
                    for k in ("supplied", "not_observed", "unknown")}
        if (self.counts != expected or self.total_entries != len(self.entries)
                or self.assessed_entries != len(self.entries) - expected["unknown"]
                or len({e.reference_id for e in self.entries}) != len(self.entries)):
            raise ValueError("Submitted locator inventory counts/bindings differ.")
        return self


def _has_locator(text: str) -> bool:
    # Presence only. Malformed/wrapped candidates must not become absences.
    return bool(_DOI_PATTERN.search(text) or _URL_PATTERN.search(text) or _LOCATOR_HINT.search(text))


def _hyperlink_fields(node):
    """Recognize complete plain HTTP fields transiently, never execute them.

    Only bookmark/window switches are recognized for presence. Other switches,
    nested fields and paragraph-crossing fields abstain.
    A target alone without a visible field result is not an actionable link.
    """
    valid = lambda instruction: len(instruction) <= 16384 and bool(re.fullmatch(
        r'\s*HYPERLINK\s+"https?://[^"\s\x00-\x1f]+"'
        r'(?:\s+\\[lt]\s+"[^"\x00-\x1f]*")*\s*', instruction, re.I))
    found = False
    for simple in node.xpath('.//w:fldSimple'):
        if (not valid(simple.get(qn('w:instr')) or '')
                or any(e.tag in {qn('w:fldSimple'), qn('w:fldChar'), qn('w:instrText')} for e in simple.iterdescendants())
                or not ''.join(t.text or '' for t in simple.iter(qn('w:t'))).strip()):
            return False, 'entry_revision_field_or_drawing_unresolved'
        found = True
    state, instruction, visible = None, '', ''
    for element in node.iter():
        if element.tag == qn('w:fldSimple') and state is not None:
            return False, 'entry_revision_field_or_drawing_unresolved'
        if element.tag == qn('w:fldChar'):
            kind = element.get(qn('w:fldCharType'))
            if kind == 'begin' and state is None:
                state, instruction, visible = 'instruction', '', ''
            elif kind == 'separate' and state == 'instruction' and valid(instruction):
                state = 'result'
            elif kind == 'end' and state == 'result' and visible.strip():
                state, found = None, True
            else:
                return False, 'entry_revision_field_or_drawing_unresolved'
        elif element.tag == qn('w:instrText'):
            if state != 'instruction':
                return False, 'entry_revision_field_or_drawing_unresolved'
            instruction += element.text or ''
            if len(instruction) > 16384:
                return False, 'entry_revision_field_or_drawing_unresolved'
        elif element.tag == qn('w:t') and state is not None:
            if state != 'result':
                return False, 'entry_revision_field_or_drawing_unresolved'
            visible += element.text or ''
    return (False, 'entry_revision_field_or_drawing_unresolved') if state else (found, None)


def _docx_entry(document, indexes):
    """Inspect existing paragraph anchors; never resolve or fetch a target."""
    channels = []
    if not indexes or len(set(indexes)) != len(indexes):
        return channels, "entry_paragraph_unavailable"
    for index in indexes:
        if index < 0 or index >= len(document.paragraphs):
            return channels, "entry_paragraph_unavailable"
        paragraph = document.paragraphs[index]
        node = paragraph._p
        if node.xpath('.//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo | .//w:drawing | .//w:pict'):
            return channels, "entry_revision_field_or_drawing_unresolved"
        field_link, limitation = _hyperlink_fields(node)
        if limitation:
            return channels, limitation
        if field_link:
            channels.append('native_hyperlink')
        text = ''.join(t.text or '' for t in node.xpath('.//w:t'))
        if not text.strip():
            return channels, "entry_text_unavailable"
        if _has_locator(text):
            channels.append("visible_text")
        for link in node.xpath('.//w:hyperlink'):
            rid = link.get(qn('r:id'))
            if not rid:
                if link.get(qn('w:anchor')):
                    continue
                return channels, "external_link_relationship_unavailable"
            relation = paragraph.part.rels.get(rid)
            if relation is None or relation.reltype != RT.HYPERLINK or not relation.is_external:
                return channels, "external_link_relationship_unavailable"
            if not _has_locator(relation.target_ref):
                return channels, "external_link_kind_uncertain"
            channels.append("native_hyperlink")
    return channels, None


def bind_submitted_hyperlinks(content: bytes, *, references: list[ParsedReference],
                             layout: ReferenceLayoutArtifact) -> list[ParsedReference]:
    """Hand a unique, original Word target to normal retrieval, not admission.

    The separate presence inventory remains text-free. Its byte/reference and
    precise-location checks are reused before opening any relationship target.
    Existing parsed locators win; ambiguous or unsupported fields abstain.
    """
    inventory = inventory_submitted_locators(content, references=references, layout=layout)
    if layout.media_type == "application/pdf":
        return references
    document = Document(io.BytesIO(content))
    entries = {entry.reference_id: entry for entry in inventory.entries}
    owners = {}
    for entry in inventory.entries:
        for index in entry.paragraph_indexes:
            owners.setdefault(index, set()).add(entry.reference_id)
    result = []
    for ref in references:
        entry = entries[ref.reference_id]
        targets = set()
        if (not ref.url and not ref.doi and entry.status == "supplied" and not entry.limitation
                and all(len(owners[index]) == 1 for index in entry.paragraph_indexes)):
            for index in entry.paragraph_indexes:
                paragraph = document.paragraphs[index]
                for link in paragraph._p.xpath('.//w:hyperlink'):
                    relation = paragraph.part.rels.get(link.get(qn('r:id')))
                    if relation is not None and relation.reltype == RT.HYPERLINK and relation.is_external:
                        targets.add(relation.target_ref)
                # Presence validation above already rejects incomplete/nested
                # fields. Concatenate split instruction runs only within this
                # independently validated paragraph.
                instructions = [n.get(qn('w:instr')) or '' for n in paragraph._p.xpath('.//w:fldSimple')]
                instructions.append(''.join(n.text or '' for n in paragraph._p.xpath('.//w:instrText')))
                for instruction in instructions:
                    targets.update(re.findall(r'HYPERLINK\s+"(https?://[^"\s]+)"', instruction, re.I))
        if len(targets) == 1:
            target = next(iter(targets))
            try:
                parsed = urlsplit(target)
                valid = (len(target) <= 8192 and parsed.scheme in {'http', 'https'}
                         and parsed.hostname and not parsed.username and not parsed.password
                         and not any(ord(c) < 32 for c in target))
            except ValueError:
                valid = False
            if valid:
                ref = ref.model_copy(update={
                    'url': target,
                    'submitted_hyperlink_binding': SubmittedHyperlinkBinding(
                        paper_sha256=inventory.paper_sha256,
                        reference_text_sha256=entry.reference_text_sha256,
                        paragraph_indexes=list(entry.paragraph_indexes)),
                })
        result.append(ref)
    return result


def inventory_submitted_locators(content: bytes, *, references: list[ParsedReference],
                                 layout: ReferenceLayoutArtifact,
                                 reference_list_coverage: Literal["complete", "partial", "unknown"] = "unknown"
                                 ) -> SubmittedLocatorInventory:
    """Inspect original PDF regions or DOCX paragraphs without fetching links.

    Unsupported formats and uncertain mappings abstain. Layout completeness is
    not proof that the whole reference list was extracted. No raw reference,
    URL, annotation target or enriched metadata is retained in this result.
    """
    if len(content) > 50_000_000 or len(references) > 2000:
        raise ValueError("Submitted inventory input exceeds its bound.")
    layout = ReferenceLayoutArtifact.model_validate(layout.model_dump())
    digest = hashlib.sha256(content).hexdigest()
    if digest != layout.content_sha256:
        raise ValueError("Submitted layout belongs to different paper bytes.")
    ids = [r.reference_id for r in references]
    if any(not key for key in ids) or len(set(ids)) != len(ids):
        raise ValueError("Submitted references need unique stable IDs.")
    mapped = {e.reference_id: e for e in layout.entries}
    if len(mapped) != len(layout.entries) or set(mapped) != set(ids):
        raise ValueError("Submitted layout/reference IDs differ.")
    for ref in references:
        if mapped[ref.reference_id].reference_text_sha256 != hashlib.sha256(ref.raw_ref.encode()).hexdigest():
            raise ValueError("Submitted reference text changed after layout extraction.")
    snapshot = [{"id": r.reference_id, "raw_ref": r.raw_ref,
                 "needs_review": r.needs_review, "extraction_method": r.extraction_method}
                for r in references]
    snapshot_hash = hashlib.sha256(json.dumps(snapshot, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    document = None
    word_document = None
    try:
        if layout.media_type == "application/pdf":
            document = fitz.open(stream=content, filetype="pdf")
        else:
            # Bound expansion before python-docx reads the package. Upload safety
            # and authorization are still the caller's separate responsibility.
            with zipfile.ZipFile(io.BytesIO(content)) as package:
                infos = package.infolist()
                if len(infos) > 10000 or sum(i.file_size for i in infos) > 100_000_000:
                    raise ValueError("Submitted DOCX expansion exceeds its bound.")
            word_document = Document(io.BytesIO(content))
        entries = []
        links_by_page = {}
        # A mapped region is not proof of a correctly segmented entry. At a
        # page transition, a continuation can have been assigned to its neighbor.
        boundary_uncertain = set()
        for left, right in zip(references, references[1:]):
            a, b = mapped[left.reference_id], mapped[right.reference_id]
            if a.rectangles and b.rectangles and max(x.page_index for x in a.rectangles) < min(x.page_index for x in b.rectangles):
                boundary_uncertain.update((left.reference_id, right.reference_id))
        for ref in references:
            item = mapped[ref.reference_id]
            channels = []
            limitation = None
            if ref.needs_review or not ref.raw_ref.strip() or item.mapping_status != "matched":
                limitation = "reference_extraction_or_mapping_uncertain"
            elif item.match_confidence < 0.95:
                limitation = "entry_mapping_not_precise_enough_for_absence"
            elif word_document is not None:
                channels, limitation = _docx_entry(word_document, item.location_indexes)
                if _has_locator(ref.raw_ref):
                    channels.append("raw_text")
            elif document is None or not item.rectangles:
                limitation = "native_hyperlinks_not_inspected"
            elif ref.reference_id in boundary_uncertain:
                limitation = "cross_page_entry_boundary_unverified"
            else:
                if _has_locator(ref.raw_ref):
                    channels.append("raw_text")
                numbered_lines = set()
                for box in item.rectangles:
                    if box.page_index >= len(document):
                        limitation = "entry_page_unavailable"; break
                    page = document[box.page_index]
                    rect = fitz.Rect(box.x0, box.y0, box.x1, box.y1)
                    if not rect.is_valid or rect.is_infinite or not page.rect.contains(rect):
                        limitation = "entry_region_unavailable"; break
                    text = page.get_text("text", clip=rect)
                    if not text.strip():
                        limitation = "entry_text_unavailable"; break
                    if _has_locator(text):
                        channels.append("visible_text")
                    if re.match(r"^\s*\d{1,3}\.\s+\D", text):
                        numbered_lines.add((box.page_index, round(box.y0, 1)))
                    if box.page_index not in links_by_page:
                        links_by_page[box.page_index] = page.get_links()
                    for link in links_by_page[box.page_index]:
                        if not rect.intersects(link["from"]):
                            continue
                        uri = link.get("uri")
                        if uri and _has_locator(uri):
                            channels.append("native_hyperlink")
                        elif link.get("kind") == fitz.LINK_URI:
                            limitation = "external_link_kind_uncertain"
                if len(numbered_lines) > 1:
                    limitation = "multiple_numbered_entries_in_mapped_region"
            status = "unknown" if limitation else "supplied" if channels else "not_observed"
            entries.append(SubmittedLocatorEntry(
                reference_id=ref.reference_id, reference_text_sha256=item.reference_text_sha256,
                status=status, evidence_channels=tuple(sorted(set(channels))),
                rectangles=tuple(item.rectangles) if item.mapping_status == "matched" else (),
                paragraph_indexes=tuple(item.location_indexes) if word_document is not None and item.mapping_status == "matched" else (),
                limitation=limitation,
            ))
    finally:
        if document is not None:
            document.close()
    counts = {k: sum(e.status == k for e in entries) for k in ("supplied", "not_observed", "unknown")}
    return SubmittedLocatorInventory(paper_sha256=digest, submitted_snapshot_sha256=snapshot_hash,
        reference_list_coverage=reference_list_coverage, entries=tuple(entries),
        total_entries=len(entries), assessed_entries=len(entries)-counts["unknown"], counts=counts)

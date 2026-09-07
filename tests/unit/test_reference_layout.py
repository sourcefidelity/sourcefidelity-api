import io

import fitz
from docx import Document
from docx.shared import Inches

from app.services.reference_layout import extract_reference_layout_from_bytes
from app.services.schemas import ParsedReference


def _reference(reference_id: str, raw_ref: str) -> ParsedReference:
    return ParsedReference(
        reference_id=reference_id,
        author=raw_ref.split(",", 1)[0],
        year="2020",
        title="Bounded title",
        raw_ref=raw_ref,
        citation_key=f"{reference_id}2020",
    )


def test_docx_reference_layout_binds_entry_and_preserves_hanging_indent():
    raw_ref = "Smith, J. (2020). Bounded title. Journal Name, 2(1), 1-9."
    document = Document()
    document.add_paragraph("References")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.add_run("Smith, J. (2020). Bounded title. ")
    italic = paragraph.add_run("Journal Name")
    italic.italic = True
    paragraph.add_run(", 2(1), 1-9.")
    content = io.BytesIO()
    document.save(content)

    result = extract_reference_layout_from_bytes(
        content.getvalue(),
        "paper.docx",
        references=[_reference("ref-1", raw_ref)],
        citation_format="apa",
    )

    assert result.status == "complete"
    assert result.heading_status == "matched"
    assert result.matched_reference_count == 1
    entry = result.entries[0]
    assert entry.mapping_status == "matched"
    assert entry.first_line_x_points == 0.0
    assert 35 <= entry.observed_hanging_indent_points <= 37
    assert entry.italic_character_fraction > 0
    assert [item.model_dump() for item in entry.text_style_spans] == [
        {
            "start": raw_ref.index("Journal Name"),
            "end": raw_ref.index("Journal Name") + len("Journal Name"),
            "italic": True,
            "bold": False,
        }
    ]
    assert raw_ref not in result.model_dump_json()


def test_pdf_reference_layout_records_rendered_lines_without_style_verdict():
    raw_ref = "Smith, J. (2020). Bounded title. Journal Name, 2(1), 1-9."
    document = fitz.open()
    page = document.new_page()
    page.insert_text((250, 72), "References", fontsize=12)
    page.insert_text((72, 110), "Smith, J. (2020). Bounded title.", fontsize=12)
    page.insert_text((108, 130), "Journal Name, 2(1), 1-9.", fontsize=12)
    content = document.tobytes()
    document.close()

    result = extract_reference_layout_from_bytes(
        content,
        "paper.pdf",
        references=[_reference("ref-1", raw_ref)],
        citation_format="apa",
    )

    assert result.status == "complete"
    entry = result.entries[0]
    assert entry.mapping_status == "matched"
    assert entry.line_count == 2
    assert 35 <= entry.observed_hanging_indent_points <= 37
    assert "no APA or MLA conformance conclusion" in result.limitations[1]


def test_reference_layout_fails_closed_without_physical_heading():
    document = Document()
    document.add_paragraph("Bibliography")
    document.add_paragraph("An unrelated paragraph without the cited record.")
    content = io.BytesIO()
    document.save(content)

    result = extract_reference_layout_from_bytes(
        content.getvalue(),
        "paper.docx",
        references=[_reference("ref-1", "Smith, J. (2020). Bounded title.")],
        citation_format="apa",
    )

    assert result.status == "not_assessed"
    assert result.heading_status == "not_matched"
    assert result.entries[0].mapping_status == "not_matched"


def test_reference_layout_can_infer_boundary_from_unique_full_first_reference():
    raw_ref = "Smith, J. (2020). Bounded title. Journal Name, 2(1), 1-9."
    document = Document()
    document.add_paragraph("A body paragraph.")
    paragraph = document.add_paragraph(raw_ref)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    content = io.BytesIO()
    document.save(content)

    result = extract_reference_layout_from_bytes(
        content.getvalue(),
        "paper.docx",
        references=[_reference("ref-1", raw_ref)],
        citation_format="apa",
    )

    assert result.status == "complete"
    assert result.heading_status == "inferred_from_first_reference"
    assert result.matched_reference_count == 1

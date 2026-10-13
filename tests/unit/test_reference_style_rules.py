from app.services.evidence_report import summary_text
import io
import pytest
from docx import Document
from app.services.parsers import detect_format
from app.services.parsers.apa_parser import ApaParser
from app.services.ref_field_extractor import extract_fields_apa
from app.services.reference_layout import extract_reference_layout_from_bytes
from app.services.reference_formatting import assess_reference_formatting, ReferenceFormattingAssessment


def build(raws, italic_titles=(), italic_journals=()):
    doc=Document();doc.add_paragraph('References');refs=[]
    for i,raw in enumerate(raws):
        ref=extract_fields_apa(raw);ref.reference_id=str(i);refs.append(ref)
        p=doc.add_paragraph();a,sep,b=raw.partition(ref.title)
        p.add_run(a);p.add_run(sep).italic=i in italic_titles
        c,journal,d=b.partition('Journal Name, 4')
        p.add_run(c);p.add_run(journal).italic=i in italic_journals;p.add_run(d)
    buf=io.BytesIO();doc.save(buf)
    layout=extract_reference_layout_from_bytes(buf.getvalue(),'case.docx',references=refs,citation_format='apa')
    return refs,layout


BOOK='Smith, J. (2020). A complete book title. University Press.'
ARTICLE='Adams, A. (2021). A complete article title. Journal Name, 4(2), 11-19.'


def test_literal_asterisks_do_not_hide_plain_book_title_or_absorb_publisher():
    raw = 'Smith, J. (2020). *A complete book title.* University Press.'
    refs, layout = build([raw])
    assert refs[0].title == 'A complete book title'
    result = assess_reference_formatting(layout, references=refs).title_results[0]
    assert result.status == 'difference'
    assert raw[result.title_start:result.title_end] == refs[0].title


def test_literal_marked_periodical_requires_actual_italics():
    raw = 'Adams, A. (2021). A complete article title. *Journal Name, 4*(2), 11-19.'
    refs, layout = build([raw])
    assessment = assess_reference_formatting(layout, references=refs)
    result = next(r for r in assessment.title_results if r.rule_id == 'apa7_marked_periodical_italics_v1')
    assert result.status == 'difference'
    assert raw[result.title_start:result.title_end] == 'Journal Name, 4'


@pytest.mark.parametrize('raw', [
    'Smith, J. (2020). Images of the “Other” in cinema. University Press.',
    'Smith, J. (2020). A complete book title. ExampleUniversityPress.',
    'Smith, J. (Director). (2020). Example film [Film]. Example Studios.',
    # APA 6's film label (paper 2's Kubrick, 2026-10-03).
    'Smith, J. (Director). (2020). Example film [Motion Picture]. Example Studios.',
])
def test_plain_book_quoted_phrase_and_explicit_film_titles(raw):
    refs, layout = build([raw])
    result = assess_reference_formatting(layout, references=refs).title_results[0]
    assert result.status == 'difference'
    assert result.rule_id == 'apa7_reference_title_italics_v2'
    assert '[' not in raw[result.title_start:result.title_end]


def test_blog_post_does_not_inherit_book_or_film_title_rule():
    refs, layout = build(['Example Blog. (2020, May 1). A blog post title. https://example.org/post'])
    assert assess_reference_formatting(layout, references=refs).title_results[0].status == 'not_assessed'


def test_apa_heading_overrides_incidental_mla_pattern_and_bullets_are_preserved():
    text='Intro\nSomeone, Example. Incidental prose.\nReferences\n- '+BOOK+'\n- '+ARTICLE
    assert detect_format(text) is ApaParser
    raws=ApaParser.split_references(text.split('References\n')[1])
    assert len(raws)==2 and raws[0].startswith('- ')
    assert extract_fields_apa(raws[0]).author == 'Smith, J'


def test_title_rule_and_separate_order_result_roundtrip():
    refs,layout=build([BOOK,ARTICLE],italic_titles=(1,))
    result=assess_reference_formatting(layout,references=refs,reference_section='\n'.join([BOOK,ARTICLE]))
    # The article's plain journal and volume are now checked too (2026-09-30).
    assert [r.status for r in result.title_results]==['difference','difference','difference']
    assert result.order_result.status=='difference'
    assert result.order_result.expected_order==['1','0']
    assert ReferenceFormattingAssessment.model_validate_json(result.model_dump_json())==result
    legacy=assess_reference_formatting(layout)
    assert not legacy.title_results and legacy.order_result is None


def test_correct_titles_and_order():
    refs,layout=build([ARTICLE,BOOK],italic_titles=(1,),italic_journals=(0,))
    result=assess_reference_formatting(layout,references=refs,reference_section=ARTICLE+'\n'+BOOK)
    assert [r.status for r in result.title_results]==['matches_rule','matches_rule','matches_rule']
    assert result.order_result.status=='matches_rule'


def test_title_can_be_assessed_without_observable_continuation_indent():
    refs,layout=build([BOOK])
    layout.entries[0].observed_hanging_indent_points=None
    result=assess_reference_formatting(layout,references=refs)
    assert result.status=='partial'
    assert result.result_counts=={'not_assessed':1}  # legacy indentation counts
    assert result.assessed_rule_ids==['apa7_reference_title_italics_v2']
    assert result.title_results[0].status=='difference'


@pytest.mark.parametrize('mutation', ['legacy','changed_hash','unknown_kind','parse_review','mixed'])
def test_title_uncertainty_never_becomes_a_difference(mutation):
    refs,layout=build([BOOK])
    if mutation=='legacy':
        layout.entries[0].style_binding_version=None;layout.entries[0].style_observation_ranges=[]
    elif mutation=='changed_hash': refs[0].raw_ref+=' changed'
    elif mutation=='unknown_kind': refs[0].source_kind='unknown'
    elif mutation=='parse_review': refs[0].needs_review=True
    else:
        from app.services.reference_layout import ReferenceTextStyleSpan
        start=BOOK.index(refs[0].title)
        layout.entries[0].text_style_spans=[ReferenceTextStyleSpan(start=start,end=start+1,italic=True)]
    result=assess_reference_formatting(layout,references=refs)
    assert result.title_results[0].status=='not_assessed'


@pytest.mark.parametrize('mode',['omitted','same_surname','merged','changed','no_lines','mla'])
def test_order_does_not_infer_completeness_from_layout(mode):
    refs,layout=build([BOOK,ARTICLE]);section=BOOK+'\n'+ARTICLE
    if mode=='omitted': section+='\nJones, B. (2019). Another book title. Press.'
    elif mode=='same_surname':
        refs,layout=build([BOOK,ARTICLE.replace('Adams','Smith')]);section='\n'.join(r.raw_ref for r in refs)
    elif mode=='merged': refs[0].raw_ref+=' '+ARTICLE;refs=refs[:1]
    elif mode=='changed': layout.entries[0].reference_text_sha256='a'*64
    elif mode=='no_lines': section=BOOK+' '+ARTICLE
    else: layout.citation_format='mla'
    result=assess_reference_formatting(layout,references=refs,reference_section=section)
    assert result.order_result.status=='not_assessed'


def test_stored_findings_bind_fields_and_suppress_unplaced_summary():
    import fitz
    import hashlib
    from app.services.reference_formatting import reference_style_findings
    from app.services.evidence_report import _reference_view, _attach_reference_field_geometry, _render_reference_panel_template, _build_report_summary
    refs,layout=build([BOOK,ARTICLE],italic_titles=(1,))
    assessment=assess_reference_formatting(layout,references=refs,reference_section=BOOK+'\n'+ARTICLE).model_dump()
    findings=reference_style_findings(assessment,{r.reference_id:r for r in refs})
    assert len(findings)==5  # includes the article's plain journal and volume
    for f in findings: f['source']=_reference_view(refs[int(f['reference_id'])])
    before=_build_report_summary(citations=[],overview={},pervasive_hanging_indent=False,reference_practice=findings,require_paper_flags=True)
    assert not any('title(s)' in s for s in map(summary_text, before['reference_formatting']))
    doc=fitz.open();page=doc.new_page()
    for i,ref in enumerate(refs): page.insert_text((72,100+40*i),ref.raw_ref,fontsize=8)
    _attach_reference_field_geometry({'reference_practice':findings},doc,hashlib.sha256(doc.tobytes()).hexdigest())
    assert all(f['localization_status']=='exact_field' for f in findings)
    assert all('APA guidance' not in _render_reference_panel_template(f,i) for i,f in enumerate(findings))
    after=_build_report_summary(citations=[],overview={},pervasive_hanging_indent=False,reference_practice=findings,require_paper_flags=True)
    assert len(after['reference_formatting'])==2
    assert any('title uses incorrect title formatting' in s or 'titles use incorrect title formatting' in s
               for s in map(summary_text, after['reference_formatting']))
    assert reference_style_findings({**assessment,'assessment_version':'reference-formatting-v1'},{r.reference_id:r for r in refs})==[]
    refs[0].raw_ref+=' changed'
    changed=reference_style_findings(assessment,{r.reference_id:r for r in refs})
    # The article keeps its title and its plain-journal findings.
    assert len(changed)==2 and {f['reference_id'] for f in changed}=={'1'}


def test_pdf_margin_number_does_not_make_next_reference_ambiguous():
    import fitz
    doc=fitz.open();page=doc.new_page()
    page.insert_text((72,72),'References')
    page.insert_text((72,100),BOOK,fontsize=9)
    page.insert_text((300,820),'1',fontsize=10)
    page=doc.new_page();page.insert_text((72,72),ARTICLE,fontsize=9)
    refs=[extract_fields_apa(BOOK),extract_fields_apa(ARTICLE)]
    for i,r in enumerate(refs): r.reference_id=str(i)
    layout=extract_reference_layout_from_bytes(doc.tobytes(),'case.pdf',references=refs,citation_format='apa')
    assert layout.matched_reference_count==2


@pytest.mark.parametrize('raw',[
    'Smith, J. (2020). A complete chapter title. In A larger book. University Press.',
    'Smith, J. (2020). The Press and public discussion.',
    'Smith, J. (2020). A complete chapter title. In A very long book title (pp. 10-20). Routledge.',
])
def test_routing_book_kind_is_not_independent_title_rule_evidence(raw):
    refs,layout=build([raw])
    refs[0].source_kind='monograph';refs[0].source_kind_confidence='high'
    result=assess_reference_formatting(layout,references=refs)
    assert result.title_results[0].status=='not_assessed'


def test_an_unmarked_journal_name_is_checked_for_italics():
    # Hess, Paper 2, 2026-09-30: a plain journal and volume were never assessed.
    refs, layout = build([ARTICLE])
    plain = next(r for r in assess_reference_formatting(layout, references=refs).title_results
                 if r.rule_id == 'apa7_marked_periodical_italics_v1')
    assert plain.status == 'difference' and ARTICLE[plain.title_start:plain.title_end] == 'Journal Name, 4'
    refs, layout = build([ARTICLE], italic_journals=(0,))
    italic = next(r for r in assess_reference_formatting(layout, references=refs).title_results
                  if r.rule_id == 'apa7_marked_periodical_italics_v1')
    assert italic.status == 'matches_rule'


def test_a_plain_part_title_before_an_italic_larger_work_is_not_flagged():
    # Curle, Academic Article, 2026-09-30.
    raw = 'Curle, S. (2020). A part title of the work. Part 1: Review. The Larger Report Title. British Council.'
    ref = extract_fields_apa(raw); ref.reference_id = '0'
    doc = Document(); doc.add_paragraph('References'); p = doc.add_paragraph()
    before, larger, after = raw.partition('The Larger Report Title')
    p.add_run(before); p.add_run(larger).italic = True; p.add_run(after)
    buf = io.BytesIO(); doc.save(buf)
    layout = extract_reference_layout_from_bytes(buf.getvalue(), 'case.docx', references=[ref], citation_format='apa')
    results = assess_reference_formatting(layout, references=[ref]).title_results
    assert all(r.status != 'difference' for r in results)


def test_a_plain_journal_and_volume_run_together_with_commas_are_checked():
    # Paper 2's Hess, 2026-10-03: "Jump Cut,1,1,16-18" was never checked.
    raw = 'Adams, A. (1974). A complete article title. Jump Cut,1,1,16-18'
    refs, layout = build([raw])
    result = next(r for r in assess_reference_formatting(layout, references=refs).title_results
                  if r.rule_id == 'apa7_marked_periodical_italics_v1')
    assert result.status == 'difference' and raw[result.title_start:result.title_end] == 'Jump Cut,1'



def test_an_edited_book_title_is_checked_for_italics():
    # Paper 7, 2026-10-04: "Schmidt, H.C. (Ed.). (2023). Issues … Taylor & Francis" went unchecked.
    raw = 'Schmidt, H. C. (Ed.). (2023). Issues in contemporary American journalism. Taylor & Francis.'
    refs, layout = build([raw])
    if refs[0].source_kind != 'edited_collection':
        pytest.skip('parser did not classify the edited book')
    result = assess_reference_formatting(layout, references=refs).title_results[0]
    assert result.status == 'difference' and result.expected_italic is True


def test_title_case_and_missing_publication_details():
    # Owner wording 2026-10-04 (paper 7).
    from app.services.reference_formatting import (PUBLICATION_MISSING_TEXT, TITLE_CASE_TEXT,
                                                   reference_entry_findings, title_case_title)
    assert title_case_title('The Yellow Press Era: Ethical Challenges and Reforms in 19th Century American Journalism')
    assert title_case_title('Freedom of the Press Under Andres Manuel López Obrador')
    assert not title_case_title('The Washington Post and the New York Times coverage of Iraq')
    assert not title_case_title('Power without responsibility: Press, broadcasting and the internet in Britain')
    assert not title_case_title('A short Title')
    refs = {}
    for i, raw in enumerate(['Johnson, L. (2021). The Yellow Press Era: Ethical Challenges and Reforms.',
                             'Adams, A. (2021). A complete article title. Journal Name, 4(2), 11-19.',
                             'Smith, J. (2020). A complete book title. University Press.']):
        ref = extract_fields_apa(raw); ref.reference_id = str(i); refs[str(i)] = ref
    found = {(f['reference_id'], f['finding']) for f in reference_entry_findings(refs)}
    assert found == {('0', TITLE_CASE_TEXT), ('0', PUBLICATION_MISSING_TEXT)}


# The mixed-case rule reads the English word lists the image installs; tests
# use a fixed list so they do not depend on the machine.
_TEST_COMMON_WORDS = frozenset('''interim provisions administration rules management online services
commentary ports commerce northern trade relations united provinces regulating everyday life power press
broadcasting may film industry speaking plainly benjamin pierce warner family corporate transcending time narrative decline actress encyclopedia'''.split())


@pytest.fixture(autouse=True)
def _english_word_list(monkeypatch):
    from app.services import reference_formatting
    monkeypatch.setattr(reference_formatting, '_common_words', lambda: _TEST_COMMON_WORDS)


def test_mixed_title_and_sentence_case_is_the_same_title_finding():
    # Owner decision 2026-10-12: counted and shown as the existing title finding.
    from app.services.reference_formatting import (MIXED_CASE_TEXT, mixed_case_title,
                                                   reference_entry_findings)
    assert mixed_case_title('Interim Rules on the Management of Internet publishing')
    assert mixed_case_title('Rules for the Management of Online services: A commentary')
    # Sentence case: "Internet" and proper nouns are not title case.
    assert not mixed_case_title('Regulating the Internet in everyday life')
    assert not mixed_case_title('Trade relations between Avonia and the United Provinces')
    assert not mixed_case_title('Trade relations between Avonia and the United Provinces in the nineteenth century')
    assert not mixed_case_title('The Washington Post and the New York Times coverage of Iraq')
    assert not mixed_case_title('Power without responsibility: Press, broadcasting and the internet in Britain')
    # A bracketed translation is ignored.
    assert not mixed_case_title('Wangluo guanli banfa (Measures on the Governance of Online information)')
    # Names a located sentence-case record capitalises do not count.
    assert mixed_case_title('Ports of Commerce and the Northern trade')
    assert not mixed_case_title('Ports of Commerce and the Northern trade',
                                ['Ports of Commerce and the Northern trade in grain'])
    # A located record in title case says nothing about names.
    assert mixed_case_title('Ports of Commerce and the Northern trade',
                            ['Ports of Commerce and the Northern Trade'])
    refs = {}
    for i, raw in enumerate(['Office, S. (2021). Interim Rules on the Management of Internet publishing. Office Press.',
                             'Adams, A. (2021). Trade relations between Avonia and the United Provinces. Journal Name, 4(2), 11-19.',
                             'Lee, B. (2020). Regulating the Internet in everyday life. University Press.']):
        ref = extract_fields_apa(raw); ref.reference_id = str(i); refs[str(i)] = ref
    found = [(f['reference_id'], f['finding'], f['rule_id']) for f in reference_entry_findings(refs)
             if f['finding_type'] == 'reference_title_style']
    assert found == [('0', MIXED_CASE_TEXT, 'apa7_title_mixed_case_v1')]
    # Names the word lists only spell capitalised never count (2026-10-12 corpus check).
    assert not mixed_case_title('Ella May Thorvald and the British film industry')
    assert not mixed_case_title("Enduring time: Jean Okafor and Hollywood's narrative of decline")
    assert not mixed_case_title('Vivien Okonkwo | British actress | Encyclopedia')
    assert not mixed_case_title('Speaking plainly: Ella May Thorvald, Walter Benjamin, and the possibilities of film')
    assert not mixed_case_title('Hildred Pierce, Warner Bros., and the corporate family')
    with pytest.MonkeyPatch.context() as patch:
        # No installed word list: the rule abstains.
        from app.services import reference_formatting
        patch.setattr(reference_formatting, '_common_words', lambda: frozenset())
        assert not mixed_case_title('Interim Rules on the Management of Internet publishing')
    refs['0'].needs_review = True
    assert not [f for f in reference_entry_findings(refs) if f['finding_type'] == 'reference_title_style']


def test_mixed_case_uses_the_located_records_capitalisation():
    from datetime import datetime, timezone
    from app.services.reference_discovery import (ExpectedBibliographicFields, ReferenceRouteAttempt,
                                                  build_reference_discovery_candidate,
                                                  derive_reference_discovery_record)
    from app.services.reference_formatting import located_titles, reference_entry_findings
    from app.services.retrieval.base import RetrievalResult
    ref = extract_fields_apa('Adams, A. (2021). Ports of Commerce and the Northern trade. Office Press.')
    ref.reference_id = 'r'

    def discovery(located):
        expected = ExpectedBibliographicFields(title=ref.title, year=ref.year, author=ref.author)
        candidate = build_reference_discovery_candidate(
            attempt_id='a', provider='crossref', expected=expected,
            result=RetrievalResult(source_name='crossref', success=True, title=located,
                                   year=ref.year, authors=['A. Adams']))
        now = datetime.now(timezone.utc)
        attempt = ReferenceRouteAttempt(attempt_id='a', provider='crossref', route_category='academic_adapter',
            required=True, permitted=True, outcome='candidate_found', started_at=now, completed_at=now)
        record = derive_reference_discovery_record(reference_id='r', expected=expected,
            required_route_categories=['academic_adapter'], queries=[], attempts=[attempt],
            candidates=[candidate])
        return record.model_dump(mode='json')

    titled = lambda d: [f for f in reference_entry_findings({'r': ref}, {'r': d})
                        if f['finding_type'] == 'reference_title_style']
    assert titled(None)
    sentence = discovery('Ports of Commerce and the Northern trade')
    assert located_titles(ref, sentence) == ['Ports of Commerce and the Northern trade']
    assert not titled(sentence)


def test_title_case_and_publication_checks_are_apa_only():
    # Owner request 2026-10-04: other styles capitalise titles differently.
    from app.services.reference_formatting import reference_style_findings
    ref = extract_fields_apa('Johnson, L. (2021). The Yellow Press Era: Ethical Challenges and Reforms.')
    ref.reference_id = 'r'
    base = {'assessment_version': 'reference-formatting-v2', 'title_results': [], 'order_result': None}
    assert reference_style_findings({**base, 'citation_format': 'apa'}, {'r': ref})
    assert reference_style_findings({**base, 'citation_format': 'mla'}, {'r': ref}) == []


def test_a_joined_reference_is_flagged_as_split():
    from app.services.reference_formatting import SPLIT_ENTRY_TEXT, reference_entry_findings
    ref = extract_fields_apa('Santillana, M. (2023). Freedom of the press under pressure. In Media (pp. 1-2). Springer.')
    ref.reference_id = 's'
    ref.split_parts = ['Santillana, M. (2023). Freedom of the press', 'under pressure. In Media (pp. 1-2). Springer.']
    assert ('s', SPLIT_ENTRY_TEXT) in {(f['reference_id'], f['finding']) for f in reference_entry_findings({'s': ref})}

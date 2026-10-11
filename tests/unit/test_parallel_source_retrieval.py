"""References searched at the same time during source retrieval (owner decision 2026-10-09)."""

import io
import threading
import time
from types import SimpleNamespace

from docx import Document
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.config import settings
from app.models import Base
from app.models.job import Job
from app.services import paper_upload
from app.services.file_safety import SafetyVerdict
from app.services.paper_upload import DOCX_MEDIA_TYPE, create_paper_job
from app.services.paper_workflow import _acquire_in_parallel, extract_paper_job, retrieve_paper_sources
from app.services.search.search_memo import DeferredSearchMemoStore, SearchMemoContext
from app.services.source_resolver import SourceResolutionError
from tests.unit.test_paper_workflow import MemoryStorage, Resolver


def _three_source_paper() -> bytes:
    document = Document()
    document.add_paragraph("Cranes lift containers (Ames, 2020).")
    document.add_paragraph("Ports automate slowly (Bell, 2021).")
    document.add_paragraph("Harbours compete on time (Cole, 2022).")
    document.add_paragraph("References")
    document.add_paragraph("Ames, J. (2020). Lifting containers. Harbour Press.")
    document.add_paragraph("Bell, A. (2021). Slow automation. Harbour Press.")
    document.add_paragraph("Cole, R. (2022). Competing on time. Harbour Press.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class SlowResolver(Resolver):
    """Records how many references were being searched at once."""

    def __init__(self, *, fail_title=None):
        super().__init__()
        self.lock = threading.Lock()
        self.active = self.peak = 0
        self.fail_title = fail_title

    def resolve_reference(self, reference):
        with self.lock:
            self.active += 1
            self.peak = max(self.peak, self.active)
        try:
            time.sleep(0.1)
            if self.fail_title and self.fail_title in (reference.title or ""):
                raise SourceResolutionError("Synthetic unavailable source")
            return super().resolve_reference(reference)
        finally:
            with self.lock:
                self.active -= 1


def _retrieve(monkeypatch, workers, resolver):
    monkeypatch.setattr(paper_upload, "scan_with_clamd", lambda _content: (SafetyVerdict.CLEAN, "OK"))
    monkeypatch.setattr(settings, "PAPER_SOURCE_RETRIEVAL_WORKERS", workers)
    engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(session, storage, content=_three_source_paper(), filename="ports.docx",
                               media_type=DOCX_MEDIA_TYPE, scope_id="personal-default")
        extract_paper_job(session, storage, job.id, llm_enabled=False)
        counts = retrieve_paper_sources(session, storage, job.id, resolver=resolver)
        job = session.get(Job, job.id)
        order = [reference["reference_id"] for reference in job.extraction_payload["references"]]
        saved = [(item["reference_id"], item["status"], item.get("reason_code")) for item in job.source_results]
    return counts, order, saved


def test_references_are_searched_at_the_same_time_and_saved_in_reference_order(monkeypatch):
    parallel = SlowResolver(fail_title="Slow automation")
    counts, order, saved = _retrieve(monkeypatch, 4, parallel)
    assert parallel.peak > 1
    one_at_a_time = SlowResolver(fail_title="Slow automation")
    serial_counts, serial_order, serial_saved = _retrieve(monkeypatch, 1, one_at_a_time)
    assert one_at_a_time.peak == 1
    assert counts == serial_counts == {"transient_authorized": 2, "unavailable": 1}
    assert [reference_id for reference_id, _, _ in saved] == order
    assert [entry[1:] for entry in saved] == [entry[1:] for entry in serial_saved]


def _plan(reference_id, doi):
    return {"reference": SimpleNamespace(reference_id=reference_id, doi=doi, title="T", raw_ref="T"),
            "record": {"reference_id": reference_id}, "identity_only": False, "reused": None, "public": None}


def test_references_to_one_work_run_one_after_another():
    spans, lock = {}, threading.Lock()

    def acquire(plan, context):
        start = time.monotonic()
        time.sleep(0.05)
        with lock:
            spans[plan["reference"].reference_id] = (start, time.monotonic(), threading.current_thread().name)
        return "result", None

    recorded = []
    plans = [_plan("a1", "10.1/same"), _plan("b", "10.1/other"), _plan("a2", "10.1/SAME")]
    _acquire_in_parallel(plans, acquire, lambda plan, result, error: recorded.append(plan["reference"].reference_id),
                         SearchMemoContext(scope_type="personal", scope_id="p"), 4)
    assert sorted(recorded) == ["a1", "a2", "b"]
    first, second = spans["a1"], spans["a2"]
    assert first[2] == second[2] and first[1] <= second[0]


def test_a_failed_search_stops_the_stage_only_after_every_worker_has_finished():
    finished = []

    def acquire(plan, context):
        if plan["reference"].reference_id == "bad":
            time.sleep(0.03)   # "slow" is under way; a search not yet started is cancelled
            raise RuntimeError("provider exploded")
        time.sleep(0.1)
        finished.append(plan["reference"].reference_id)
        return "result", None

    plans = [_plan("bad", "10.1/a"), _plan("slow", "10.1/b")]
    with pytest.raises(RuntimeError):
        _acquire_in_parallel(plans, acquire, lambda *args: None,
                             SearchMemoContext(scope_type="personal", scope_id="p"), 4)
    assert finished == ["slow"]
    assert not [thread for thread in threading.enumerate() if thread.name.startswith("source-retrieval")]


class _Store:
    def __init__(self):
        self.calls = []

    def get(self, *args):
        return SimpleNamespace(id="m", memo_policy_version="v", search_policy_version="s",
                               outcome_summary={"a": 1}, created_at=None, expires_at=None, extra="kept out")

    def put(self, *args, **kwargs):
        self.calls.append(("put", args))

    def delete(self, *args, **kwargs):
        raise RuntimeError("table missing")


def test_memo_writes_wait_for_the_reference_checkpoint():
    base = _Store()
    store = DeferredSearchMemoStore(base, threading.Lock())
    snapshot = store.get("personal", "p", "k")
    assert snapshot.outcome_summary == {"a": 1} and not hasattr(snapshot, "extra")
    store.put("personal", "p", "k", {}, created_at=None, expires_at=None)
    store.delete("personal", "p", "k")
    assert base.calls == []
    store.apply()   # a failed memo write never fails the run
    assert base.calls == [("put", ("personal", "p", "k", {}))]
    assert store.pending == []


def test_a_reference_never_reports_the_previous_references_doi_mismatch(monkeypatch):
    from app.services import source_resolver as module
    from app.services.retrieval.base import RetrievalResult
    from app.services.schemas import ParsedReference

    resolver = object.__new__(module.SourceResolver)
    discovery = {}
    monkeypatch.setattr(module.SourceResolver, "_discovery_artifacts",
                        staticmethod(lambda: ({}, discovery.setdefault("current", {}))))
    for name in ("_enrich_journal_registration", "_attach_container_identity"):
        monkeypatch.setattr(module.SourceResolver, name, lambda self, *args: None)
    monkeypatch.setattr(module.SourceResolver, "_enrich_book_editions", lambda self: None)
    monkeypatch.setattr(module.SourceResolver, "_update_search_memo", lambda self, *args: None)

    def resolve(self, **fields):
        if fields["title"] == "First":
            # The first reference's search rejects a DOI naming another work.
            module._LAST_GRAPH.set(SimpleNamespace(rejected_candidates=[{
                "rejection_code": "doi_registers_a_different_title", "doi": "10.1/other",
                "title": "Other", "provider": "crossref"}]))
        # The second returns before any graph is built.
        return RetrievalResult(source_name="test", success=True, full_text=b"text")

    monkeypatch.setattr(module.SourceResolver, "resolve", resolve)

    def reference(title):
        return ParsedReference(reference_id=title, title=title, author="Ames", year="2020",
                               source_kind="journal_article", raw_ref=f"Ames, J. (2020). {title}.")

    first = resolver.resolve_reference(reference("First"))
    assert first.metadata["reference_discovery"]["doi_title_mismatch"]["doi"] == "10.1/other"
    discovery.clear()
    second = resolver.resolve_reference(reference("Second"))
    assert "doi_title_mismatch" not in second.metadata["reference_discovery"]

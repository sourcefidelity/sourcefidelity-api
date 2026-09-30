from app.services.evidence_report import render_evidence_report_html, _locator_inventory_view
from types import SimpleNamespace
import io
import hashlib
import fitz
from docx import Document
from app.services.schemas import ParsedReference
from app.services.reference_layout import extract_reference_layout_from_bytes
from app.services.submitted_locator_inventory import inventory_submitted_locators


def inventory_view():
    return {'title':'Inventory preview', 'citation_format':'APA', 'citations':[],
        'paper_surface':{'page_dimensions':[{'page_index':0,'width':612,'height':792}],
            'page_href_template':'/page-{page_index}.png'},
        'submitted_locator_inventory':{'counts':{'supplied':1,'not_observed':1,'unknown':1},
            'assessed_entries':2, 'reference_list_coverage':'unknown', 'entries':[
                {'status':'supplied','submitted_reference':'With link','rectangles':[]},
                {'status':'not_observed','submitted_reference':'Smith, J. (2020). Book. Publisher.',
                 'rectangles':[{'page_index':0,'x0':72,'y0':100,'x1':480,'y1':125}]},
                {'status':'unknown','submitted_reference':'<script>unsafe</script>', 'rectangles':[]}]}}


def test_neutral_count_only_in_formatting_summary():
    for audience in ('student','instructor'):
        html = render_evidence_report_html({**inventory_view(),'audience':audience}, csp_nonce='inventory-test-nonce')
        assert '1 of 2 assessed references contain neither a DOI nor a URL.' not in html
        assert '1 could not be assessed.' not in html
        assert 'This does not establish that a link is required' not in html
        assert 'Smith, J.' not in html and 'unsafe' not in html
        assert 'locator-reference-' not in html
        assert 'Show reference on page' not in html
        assert 'class="locator-count"' not in html


def test_historical_inventory_not_invented():
    assert _locator_inventory_view(SimpleNamespace(submitted_locator_inventory=None, reference_layout=None)) is None
    view=inventory_view(); view.pop('submitted_locator_inventory')
    html=render_evidence_report_html(view,csp_nonce='inventory-test-nonce')
    assert 'class="locator-count"' not in html


def test_docx_navigation_requires_both_original_and_presentation_bindings():
    raw='Smith, J. (2020). Book. Publisher.'
    doc=Document();doc.add_paragraph('References');doc.add_paragraph(raw)
    stream=io.BytesIO();doc.save(stream);content=stream.getvalue()
    refs=[ParsedReference(reference_id='r',raw_ref=raw)]
    original=extract_reference_layout_from_bytes(content,'paper.docx',references=refs,citation_format='apa')
    inv=inventory_submitted_locators(content,references=refs,layout=original)
    with fitz.open() as pdf:
        page=pdf.new_page();page.insert_text((72,72),'References');page.insert_text((72,110),raw)
        rendered_bytes=pdf.tobytes()
    rendered=extract_reference_layout_from_bytes(rendered_bytes,'paper.pdf',references=refs,citation_format='apa')
    extraction=SimpleNamespace(submitted_locator_inventory=inv,reference_layout=original,references=refs)
    surface={'presentation_sha256':hashlib.sha256(rendered_bytes).hexdigest(),
        'submitted_reference_navigation':{'input_sha256':inv.paper_sha256,'layout':rendered.model_dump(mode='json')}}
    value=_locator_inventory_view(extraction,surface)
    assert value['entries'][0]['rectangles']
    assert value['counts']==inv.counts
    assert not inv.entries[0].rectangles
    for modified in ({**surface,'presentation_sha256':'0'*64},
                     {**surface,'submitted_reference_navigation':{**surface['submitted_reference_navigation'],'input_sha256':'0'*64}}):
        assert not _locator_inventory_view(extraction,modified)['entries'][0]['rectangles']

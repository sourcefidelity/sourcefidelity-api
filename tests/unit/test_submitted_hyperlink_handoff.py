import hashlib
import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.services.schemas import ParsedReference
from app.services.reference_layout import extract_reference_layout_from_bytes
from app.services.submitted_locator_inventory import bind_submitted_hyperlinks


def fixture(targets):
    document = Document()
    document.add_paragraph('References')
    raw = 'Smith, J. (2020). A bounded title. Publisher.'
    paragraph = document.add_paragraph(raw)
    for target in targets:
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
        paragraph._p.append(link)
    output = io.BytesIO()
    document.save(output)
    content = output.getvalue()
    refs = [ParsedReference(reference_id='r1', raw_ref=raw)]
    layout = extract_reference_layout_from_bytes(content, 'paper.docx', references=refs, citation_format='apa')
    return content, refs, layout


def test_unique_target_has_original_provenance_without_rewriting_reference():
    content, refs, layout = fixture(['https://example.org/source'])
    result = bind_submitted_hyperlinks(content, references=refs, layout=layout)
    assert result[0].url == 'https://example.org/source'
    assert result[0].raw_ref == refs[0].raw_ref
    assert not refs[0].url
    assert result[0].submitted_hyperlink_binding.paper_sha256 == hashlib.sha256(content).hexdigest()
    assert result[0].submitted_hyperlink_binding.paragraph_indexes == [1]


@pytest.mark.parametrize('targets', [[], ['https://example.org/a', 'https://example.org/b'], ['https://user:secret@example.org/a']])
def test_ambiguous_or_credential_target_abstains(targets):
    content, refs, layout = fixture(targets)
    assert not bind_submitted_hyperlinks(content, references=refs, layout=layout)[0].url


def test_existing_locator_and_uncertain_mapping_preserved():
    content, refs, layout = fixture(['https://example.org/source'])
    refs[0].url = 'https://example.org/original'
    assert bind_submitted_hyperlinks(content, references=refs, layout=layout)[0].url == refs[0].url
    refs[0].url = ''
    refs[0].needs_review = True
    assert not bind_submitted_hyperlinks(content, references=refs, layout=layout)[0].url
    with pytest.raises(ValueError):
        bind_submitted_hyperlinks(content+b'changed', references=refs, layout=layout)


def test_field_target_and_duplicate_relationship_share_one_destination():
    content, refs, layout = fixture(['https://example.org/source'])
    document = Document(io.BytesIO(content))
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'HYPERLINK "https://example.org/source"')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = ' source'
    run.append(text)
    field.append(run)
    document.paragraphs[1]._p.append(field)
    output = io.BytesIO()
    document.save(output)
    content = output.getvalue()
    refs[0].raw_ref = document.paragraphs[1].text
    layout = extract_reference_layout_from_bytes(content, 'paper.docx', references=refs, citation_format='apa')
    assert bind_submitted_hyperlinks(content, references=refs, layout=layout)[0].url == 'https://example.org/source'


def test_doi_is_not_overwritten_by_hidden_target():
    content, refs, layout = fixture(['https://example.org/source'])
    refs[0].doi = '10.1234/retained'
    result = bind_submitted_hyperlinks(content, references=refs, layout=layout)[0]
    assert result.doi == refs[0].doi and not result.url

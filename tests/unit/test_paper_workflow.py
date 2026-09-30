"""End-to-end local checkpoint test for the shadow paper workflow."""

import io
import json
from datetime import datetime, timezone
import hashlib
from types import SimpleNamespace
import uuid

from docx import Document
import fitz
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.models import Base
from app.models.job import Job
from app.models.report import Report, ReportPaperArtifactRecord, VerificationReportRecord
from app.models.verification_run import VerificationRunRecord
from app.services import paper_upload
from app.services.file_safety import SafetyVerdict
from app.services.evidence_report import (
    EvidenceReportAuthorizationError,
    _eligible_display_passages,
    _passage_view,
    _quotation_difference_diagnostics,
    _quotation_difference_label,
    _responsive_display_excerpt,
    _render_panel_template,
    _render_member,
    _unavailable_member,
    _prioritize_display_passages,
    get_authorized_evidence_report_view,
    load_authorized_evidence_report_bundle,
    render_evidence_report_html,
)
from app.services.paper_upload import DOCX_MEDIA_TYPE, create_paper_job
from app.services.report_paper_artifact import (
    ReportPaperArtifactError,
    load_authorized_report_paper_artifact,
)
from app.services.paper_workflow import (
    _shadow_artifact,
    _claims_by_reference,
    _report_source_failure_reason,
    _retryable_provider_dependencies,
    extract_paper_job,
    finalize_paper_job,
    prepare_provider_recovery_refresh,
    prepare_uploaded_source_refresh,
    retrieve_paper_sources,
    rollback_targeted_source_refresh,
    verify_paper_sources,
)
from app.services.verification_evidence import (
    AuthorizedRepresentation,
    ClaimEvidence,
    PassageRelevanceGateEvidence,
)
from app.services.retrieval.base import (
    RepresentationKind,
    RetrievalResult,
    SourceRepresentation,
)
from app.services.storage.backend import StorageBackend
from app.services.source_repository import (
    AdmissionRequest,
    WorkIdentity,
    admit_representation,
    commit_source_admissions,
)


class MemoryStorage(StorageBackend):
    def __init__(self):
        self.objects = {}

    def upload(self, file_bytes, key):
        self.objects[key] = file_bytes
        return key

    def download(self, key):
        if key not in self.objects:
            raise FileNotFoundError(key)
        return self.objects[key]

    def delete(self, key):
        self.objects.pop(key, None)
        return True

    def exists(self, key):
        return key in self.objects

    def list_keys(self, prefix):
        return [key for key in self.objects if key.startswith(prefix)]


@pytest.mark.parametrize("judgments_enabled", [False, True])
@pytest.mark.parametrize("joint_enabled", [False, True])
def test_shadow_workflow_runs_semantic_rescue_only_after_confirmed_miss(monkeypatch, judgments_enabled, joint_enabled):
    content = (
        b"A complete source discusses careful evidence retrieval and verification."
    )
    now = datetime.now(timezone.utc)
    source = AuthorizedRepresentation(
        representation_id="workflow-semantic-source",
        canonical_work_id="workflow-semantic-work",
        content_object_id="workflow-semantic-object",
        content_sha256=hashlib.sha256(content).hexdigest(),
        content=content,
        representation_kind="plain_text",
        media_type="text/plain",
        provenance="authorized_upload",
        scope_type="personal_owner",
        scope_id="owner-1",
        identity_verdict="verified",
        identity_confidence=0.99,
        completeness_verdict="complete",
        text_quality="digital",
        edition_or_version=None,
        created_at=now,
        admitted_at=now,
    )
    text = "Careful retrieval improves verification (Smith, 2020)."
    claim = ClaimEvidence(
        claim_id="workflow-semantic-claim",
        paper_version_id="paper-v1",
        text=text,
        reference_ids=["ref-1"],
        citation_marker="(Smith, 2020)",
        citation_marker_type="parenthetical",
        passage_start=0,
        passage_end=len(text),
    )
    calls = []

    def gate(artifact, **_kwargs):
        assert _kwargs['source_title'] == 'Submitted study title'
        calls.append("gate")
        outcome = (
            "no_relevant_candidate_passage"
            if calls.count("gate") == 1
            else "relevant_candidates_found"
        )
        return artifact.model_copy(
            update={
                "passage_relevance": PassageRelevanceGateEvidence(
                    status="complete",
                    method="test_gate",
                    outcome=outcome,
                )
            }
        )

    def interpretations(artifact, **_kwargs):
        calls.append("interpretation")
        return artifact

    def rescue(_source, artifact):
        calls.append("rescue")
        retrieval = artifact.candidate_passage_retrieval.model_copy(
            update={
                "semantic_rescue_status": "complete",
                "semantic_addition_count": 1,
            }
        )
        return artifact.model_copy(
            update={"candidate_passage_retrieval": retrieval}
        )

    def foundation(artifact):
        calls.append("foundation")
        return artifact

    def judgment(artifact):
        calls.append("judgment")
        return artifact

    def critic(artifact):
        calls.append("critic")
        return artifact

    def joint(artifact, *, source_title):
        assert source_title == 'Submitted study title'
        assert artifact.passage_relevance.outcome == 'relevant_candidates_found'
        calls.append('joint')
        return artifact.joint_evidence_selection

    monkeypatch.setattr('app.services.joint_evidence_selection.select_joint_evidence', joint)
    def selection(artifact):
        calls.append('selection')
        return artifact
    monkeypatch.setattr('app.services.evidence_sentence_selection.attach_sentence_evidence', selection)
    monkeypatch.setattr('app.services.paper_workflow.settings.PAPER_EXPERIMENTAL_JOINT_SELECTION_ENABLED', joint_enabled)

    monkeypatch.setattr(
        "app.services.paper_workflow.settings.EVIDENCE_RETRIEVAL_SEMANTIC_BACKEND",
        "deberta_nli",
    )
    monkeypatch.setattr(
        "app.services.paper_workflow.attach_source_blind_interpretations",
        interpretations,
    )
    monkeypatch.setattr("app.services.paper_workflow.apply_passage_relevance_gate", gate)
    monkeypatch.setattr(
        "app.services.paper_workflow.attach_local_semantic_retrieval_rescue", rescue
    )
    monkeypatch.setattr(
        "app.services.paper_workflow.attach_facet_evidence_foundation", foundation
    )
    monkeypatch.setattr(
        "app.services.paper_workflow.apply_facet_evidence_judgment", judgment
    )
    monkeypatch.setattr("app.services.paper_workflow.apply_decisive_label_critic", critic)
    monkeypatch.setattr("app.services.paper_workflow.settings.PAPER_EXPERIMENTAL_RELATIONSHIP_JUDGMENTS_ENABLED", judgments_enabled)

    artifact = _shadow_artifact(
        source,
        claim,
        True,
        active_reference_id="ref-1",
        cited_author_label="Smith",
        source_title='Submitted study title',
    )

    assert calls == [
        "interpretation",
        "gate",
        "rescue",
        "gate",
    ] + (["joint"] if joint_enabled else []) + ["selection", "foundation"] + (["judgment", "critic"] if judgments_enabled else [])
    assert artifact.passage_relevance.outcome == "relevant_candidates_found"
    assert artifact.candidate_passage_retrieval.semantic_rescue_status == "complete"
    calls.clear()
    _shadow_artifact(source, claim, False, active_reference_id='ref-1',
                     cited_author_label='Smith', source_title='Submitted study title')
    assert calls == ['foundation']


def test_authorized_report_bundle_requires_hash_bound_pdf_surface(monkeypatch):
    artifact = SimpleNamespace(
        id="artifact-1",
        presentation_status="page_faithful_ready",
        presentation_media_type="application/pdf",
    )
    view = {
        "paper_surface": {
            "artifact_id": "artifact-1",
            "anchor_reason_code": "presentation_hash_bound",
        }
    }
    monkeypatch.setattr(
        "app.services.evidence_report.get_authorized_evidence_report_view",
        lambda *args, **kwargs: view,
    )
    monkeypatch.setattr(
        "app.services.evidence_report.load_authorized_report_paper_artifact",
        lambda *args, **kwargs: (artifact, b"%PDF-bound"),
    )

    loaded_view, loaded_artifact, content = load_authorized_evidence_report_bundle(
        object(),
        object(),
        report_id="00000000-0000-0000-0000-000000000001",
        artifact_id="00000000-0000-0000-0000-000000000002",
        scope_type="personal_owner",
        scope_id="owner-1",
    )

    assert loaded_view is view
    assert loaded_artifact is artifact
    assert content == b"%PDF-bound"


def _assert_only_report_marking_copy(storage):
    assert len(storage.objects) == 1
    assert next(iter(storage.objects)).startswith("report-paper-artifacts/")


class Resolver:
    def __init__(self):
        self.calls = 0

    def resolve_reference(self, reference):
        self.calls += 1
        return RetrievalResult(
            source_name="test_source",
            success=True,
            representation=SourceRepresentation(
                kind=RepresentationKind.PLAIN_TEXT,
                media_type="text/plain",
                content=(
                    b"Careful verification improves accuracy by checking every source. "
                    b"The process records inspectable evidence."
                ),
                completeness="complete",
            ),
            title=reference.title,
            metadata={
                "identity_confidence": "high",
                "text_quality": "digital",
                "edition_or_version": "author_accepted_manuscript",
                "reference_discovery_trace": {
                    "trace_version": "reference-discovery-trace-v1",
                    "outcome_derived": True,
                },
                "reference_discovery": {
                    "record_version": "reference-discovery-v1",
                    "outcome": "confirmed",
                },
            },
        )


class ProvisionalResolver(Resolver):
    def resolve_reference(self, reference):
        result = super().resolve_reference(reference)
        result.metadata['identity_confidence'] = 'medium'
        result.metadata['provisional_source'] = {
            'policy_version': 'submission-possible-match-v1',
            'content_sha256': hashlib.sha256(result.full_text).hexdigest(),
            'reason': 'The publication year differs from the reference.',
        }
        return result


def test_possible_match_paper_pipeline_cleans_source_and_preserves_label(monkeypatch):
    from app.models.source_repository import SourceRepresentationRecord
    monkeypatch.setattr(paper_upload, 'scan_with_clamd',
        lambda _: (SafetyVerdict.CLEAN, 'stream: OK'))
    engine = create_engine('sqlite+pysqlite:///:memory:')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(session, storage, content=_paper_bytes(),
            filename='paper.docx', media_type=DOCX_MEDIA_TYPE, scope_id='personal-default')
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        assert retrieve_paper_sources(session, storage, job_id,
            resolver=ProvisionalResolver()) == {'transient_authorized': 1}
        assert session.scalar(select(VerificationRunRecord)).identity_verdict == 'possible_match'
        assert session.scalar(select(SourceRepresentationRecord)) is None
    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert summary['reports_persisted'] == 2
    with factory() as session:
        assert session.scalar(select(VerificationRunRecord)).cleaned_at is not None
        records = session.scalars(select(VerificationReportRecord)).all()
        for record in records:
            assert record.verdict == 'not_assessed'
            assert record.report_payload['source_identity']['status'] == 'uncertain'
            assert 'The publication year differs from the reference.' in record.report_payload['source_identity']['limitations']
        finalize_paper_job(session, storage, job_id)
        aggregate = session.scalar(select(Report))
        view = aggregate.report_json['evidence_report']
        html = render_evidence_report_html(view, csp_nonce='test-provisional-nonce')
        assert 'Possible source match' in html
        assert 'The publication year differs from the reference.' in html
        for citation in view['citations']:
            for member in citation['members']:
                assert member['identity_status'] == 'possible_match'
                assert not member['show_quotation_check']
                assert not member['show_locator_check']
                assert member['relevance_status'] == 'not_assessed'
        assert not any(key.startswith('verification-runs/') for key in storage.objects)


class CompoundResolver:
    def __init__(self):
        self.calls = []

    def resolve_reference(self, reference):
        self.calls.append(reference.reference_id)
        surname = (reference.author or "").split(",", 1)[0]
        source_text = (
            f"The {surname} source independently establishes that the shared result "
            "is stable through repeated trials."
        )
        return RetrievalResult(
            source_name="compound_test_source",
            success=True,
            representation=SourceRepresentation(
                kind=RepresentationKind.PLAIN_TEXT,
                media_type="text/plain",
                content=source_text.encode(),
                completeness="complete",
            ),
            title=reference.title,
            authors=[reference.author],
            year=reference.year,
            metadata={
                "identity_confidence": "high",
                "text_quality": "digital",
            },
        )


class PartialCompoundResolver(CompoundResolver):
    def resolve_reference(self, reference):
        if (reference.author or "").startswith("Jones"):
            self.calls.append(reference.reference_id)
            return RetrievalResult(
                source_name="compound_test_source",
                success=False,
                title=reference.title,
                authors=[reference.author],
                year=reference.year,
                error="not available",
            )
        return super().resolve_reference(reference)


def _paper_bytes():
    document = Document()
    document.add_paragraph("Careful verification improves accuracy (Smith, 2020).")
    document.add_paragraph("Inspectable evidence improves review (Smith, 2020).")
    document.add_paragraph("References")
    document.add_paragraph("Smith, J. (2020). A useful title. Example Press.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _compound_paper_bytes(*, with_link=False):
    document = Document()
    document.add_paragraph(
        "Smith (2020) and Jones (2021) show that the shared result is stable."
    )
    document.add_paragraph("References")
    document.add_paragraph("Smith, J. (2020). First source. Example Press.")
    document.add_paragraph("Jones, A. (2021). Second source. Example Press." + (' https://example.org/source' if with_link else ''))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _retry_test_job(monkeypatch, *, store_only=False, factory=None, storage=None):
    monkeypatch.setattr(paper_upload, "scan_with_clamd", lambda _content: (SafetyVerdict.CLEAN, "OK"))
    if factory is None:
        engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool)
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine, expire_on_commit=False)
    if storage is None:
        storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(session, storage, content=_paper_bytes(), filename="retry.docx",
                               media_type=DOCX_MEDIA_TYPE, scope_id="personal-default", store_only=store_only)
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        if not store_only:
            retrieve_paper_sources(session, storage, job_id, resolver=Resolver())
    return factory, storage, job_id


def test_uncited_identity_checkpoint_never_creates_source_run(monkeypatch):
    from app.services.bibliography_identity import POLICY
    from app.services import paper_workflow
    factory, storage, job_id = _retry_test_job(monkeypatch)
    calls = []
    class IdentityResolver:
        def prefetch_title_candidates(self, queries):
            calls.append(("prefetch", len(queries)))
        def prefetch_deferred_dois(self, dois):
            pass
        def resolve_reference(self, reference, *, identity_only=False):
            assert identity_only and reference.reference_id == "uncited-book"
            calls.append(("identity", reference.reference_id))
            # Even a defective injected result cannot authorize source use.
            return RetrievalResult(source_name="test", success=True,
                full_text=b"Must not store these bytes", abstract="Must not use this abstract",
                metadata={"repository_representation_id": str(uuid.uuid4()),
                          "identity_confidence": "high"})
    with factory() as session:
        job = session.get(Job, job_id)
        original_sources = list(job.source_results)
        payload = dict(job.extraction_payload)
        payload["references"] = [*payload["references"],
            dict(payload["references"][0], reference_id="uncited-book", source_kind="monograph")]
        job.extraction_payload = payload
        job.stage = "extracted"
        session.commit()
        objects = dict(storage.objects)
        runs = list(session.scalars(select(VerificationRunRecord.id)))
        retrieve_paper_sources(session, storage, job_id, resolver=IdentityResolver())
        saved = job.source_results[-1]
        assert saved["identity_only_policy"] == POLICY
        assert saved["status"] == "metadata_only"
        assert not {"representation_id", "run_id", "abstract_evidence"} & saved.keys()
        assert job.source_results[:-1] == original_sources
        assert storage.objects == objects
        assert list(session.scalars(select(VerificationRunRecord.id))) == runs
        before = list(calls)
        retrieve_paper_sources(session, storage, job_id, resolver=IdentityResolver())
        assert calls == before


def test_cited_sources_keep_priority_over_uncited_search_budget(monkeypatch):
    from app.services.source_resolver import SourceResolutionError
    factory, storage, job_id = _retry_test_job(monkeypatch)
    calls = []
    class NoSources:
        def resolve_reference(self, reference, *, identity_only=False):
            calls.append((reference.reference_id, identity_only))
            raise SourceResolutionError("Synthetic unavailable source")
    with factory() as session:
        job = session.get(Job, job_id)
        payload = dict(job.extraction_payload)
        cited_id = payload["references"][0]["reference_id"]
        payload["references"] = [dict(payload["references"][0],
            reference_id="alphabetically-first-uncited", source_kind="monograph"), *payload["references"]]
        job.extraction_payload = payload
        job.source_results = []
        job.stage = "extracted"
        session.commit()
        retrieve_paper_sources(session, storage, job_id, resolver=NoSources())
        assert calls == [(cited_id, False), ("alphabetically-first-uncited", True)]


@pytest.mark.parametrize("store_only", [False, True])
def test_finalization_resumes_after_projection_failure(monkeypatch, store_only):
    from app.services import paper_workflow
    factory, storage, job_id = _retry_test_job(monkeypatch, store_only=store_only)
    if not store_only:
        verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    original = paper_workflow.build_evidence_report_view
    def fail_once(**kwargs):
        monkeypatch.setattr(paper_workflow, "build_evidence_report_view", original)
        raise ConnectionError("temporary failure")
    monkeypatch.setattr(paper_workflow, "build_evidence_report_view", fail_once)
    with pytest.raises(ConnectionError), factory() as session:
        finalize_paper_job(session, storage, job_id)
    with factory() as session:
        assert session.get(Job, job_id).stage == "finalizing"
        assert session.scalars(select(Report)).all() == []
        result = finalize_paper_job(session, storage, job_id)
        assert finalize_paper_job(session, storage, job_id)["report_id"] == result["report_id"]
        assert len(session.scalars(select(Report)).all()) == 1


class CandidateAuditResolver(Resolver):
    """Discovery records carrying the development-only candidate-link audit."""

    AUDIT_URL = "https://exa.example/CANDIDATE_AUDIT_URL_SENTINEL.pdf"

    def resolve_reference(self, reference):
        result = super().resolve_reference(reference)
        audit = {"audit_policy": "dev-candidate-audit-v1", "omitted_count": 0, "candidates": [
            {"provider": "exa", "query_id": "query-audit", "rank": 1, "url": self.AUDIT_URL,
             "url_truncated": False, "disposition": "identity_rejected",
             "reason_code": "identity_rejected"}]}
        attempt = {"attempt_id": "attempt-audit", "route_category": "bounded_web",
                   "provider": "web_search", "candidate_audit": audit}
        result.metadata["reference_discovery"]["attempts"] = [attempt]
        result.metadata["reference_discovery_trace"]["attempts"] = [dict(attempt)]
        return result


def test_report_projection_never_renders_candidate_audit_urls(monkeypatch):
    monkeypatch.setattr(paper_upload, "scan_with_clamd", lambda _content: (SafetyVerdict.CLEAN, "OK"))
    engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    sentinel = "CANDIDATE_AUDIT_URL_SENTINEL"
    with factory() as session:
        job = create_paper_job(session, storage, content=_paper_bytes(), filename="audit.docx",
                               media_type=DOCX_MEDIA_TYPE, scope_id="personal-default")
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        retrieve_paper_sources(session, storage, job_id, resolver=CandidateAuditResolver())
        # The development record is present in the stored source results...
        assert sentinel in json.dumps(session.get(Job, job_id).source_results)
    verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        finalize_paper_job(session, storage, job_id)
        aggregate = session.scalar(select(Report))
        view = aggregate.report_json["evidence_report"]
        html = render_evidence_report_html(view, csp_nonce="test-candidate-audit-nonce")
        payloads = [record.report_payload for record in session.scalars(select(VerificationReportRecord))]
        # ...but no report projection, rendering or package carries it.
        assert sentinel not in json.dumps(aggregate.report_json, default=str)
        assert sentinel not in html
        assert "candidate_audit" not in json.dumps(aggregate.report_json, default=str)
        assert payloads and sentinel not in json.dumps(payloads, default=str)


@pytest.mark.parametrize("boundary", ["processing", "partial_batch"])
def test_verification_retry_keeps_source_before_batch_commit(monkeypatch, boundary):
    from app.services import paper_workflow, verification_run
    factory, storage, job_id = _retry_test_job(monkeypatch)
    module = paper_workflow if boundary == "processing" else verification_run
    name = "_shadow_artifact" if boundary == "processing" else "persist_verification_report"
    original = getattr(module, name)
    calls = 0

    def interrupted(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ConnectionError("temporary synthetic interruption")
        return original(*args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(module, name, interrupted)
        with pytest.raises(ConnectionError):
            verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        run = session.scalar(select(VerificationRunRecord))
        assert run.status == "active"
        assert run.terminal_outcome is None
        assert run.transient_objects
        assert all(storage.exists(item["key"]) for item in run.transient_objects)
        assert session.scalars(select(VerificationReportRecord)).all() == []
        assert session.get(Job, job_id).verification_summary is None
    result = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert result["reports_persisted"] == 2
    assert result["source_failures"] == []
    with factory() as session:
        assert session.scalar(select(VerificationRunRecord)).status == "cleaned"
        assert len(session.scalars(select(VerificationReportRecord)).all()) == 2


@pytest.mark.parametrize("error_type,starts", [(ValueError, 1), (ConnectionError, 4)])
def test_verification_nonretryable_or_exhausted_attempt_cleans(monkeypatch, error_type, starts):
    from app.services import paper_workflow
    from app.services.paper_dispatch import DISPATCH_KEY
    factory, storage, job_id = _retry_test_job(monkeypatch)
    with factory() as session:
        job = session.get(Job, job_id)
        job.upload_evidence = {**job.upload_evidence, DISPATCH_KEY: {"stage_starts": {"verify": starts}}}
        session.commit()

    def fail(*args, **kwargs):
        raise error_type("synthetic processing failure")

    monkeypatch.setattr(paper_workflow, "_shadow_artifact", fail)
    with pytest.raises(error_type):
        verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        assert session.scalar(select(VerificationRunRecord)).status == "failed_cleaned"
        assert session.scalars(select(VerificationReportRecord)).all() == []
        assert session.get(Job, job_id).verification_summary is None
    assert storage.list_keys("verification-runs/") == []


@pytest.mark.parametrize("damage", ["cleaned", "tampered", "expired", "missing"])
def test_interrupted_transient_source_never_becomes_successful_missing_evidence(monkeypatch, damage):
    from datetime import timedelta
    from app.services.paper_workflow import PaperWorkflowError
    from app.services.verification_run import cleanup_verification_run
    factory, storage, job_id = _retry_test_job(monkeypatch)
    with factory() as session:
        run = session.scalar(select(VerificationRunRecord))
        key = run.transient_objects[0]["key"]
        if damage == "cleaned":
            cleanup_verification_run(session, storage, run.id, outcome="failed")
        elif damage == "expired":
            run.lease_expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
            session.commit()
        elif damage == "missing":
            storage.delete(key)
        else:
            storage.upload(b"Different source bytes", key)
    with pytest.raises((PaperWorkflowError, FileNotFoundError)):
        verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        assert session.get(Job, job_id).verification_summary is None
        assert session.get(Job, job_id).stage == "verifying"
        assert session.scalars(select(VerificationReportRecord)).all() == []


def test_verification_retry_recovers_a_lost_batch_commit_acknowledgement(monkeypatch):
    from sqlalchemy import event
    from app.services import paper_workflow
    from app.services.verification_run import cleanup_stale_verification_runs
    factory, storage, job_id = _retry_test_job(monkeypatch)

    def lost_ack(session):
        if any(isinstance(row, VerificationRunRecord) and row.status == "report_persisted"
               for row in session.identity_map.values()):
            raise ConnectionError("synthetic lost batch commit acknowledgement")

    event.listen(factory.class_, "after_commit", lost_ack)
    try:
        with pytest.raises(ConnectionError):
            verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    finally:
        event.remove(factory.class_, "after_commit", lost_ack)
    with factory() as session:
        saved = {str(row.id): row.evidence_sha256 for row in session.scalars(select(VerificationReportRecord))}
        assert len(saved) == 2
        assert session.scalar(select(VerificationRunRecord)).status == "report_persisted"
    monkeypatch.setattr(paper_workflow, "_shadow_artifact", lambda *a, **kw: pytest.fail("committed evidence must be reused"))
    monkeypatch.setattr(storage, "download", lambda key: pytest.fail("source must not be reopened"))
    result = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert set(result["report_ids"]) == set(saved)
    with factory() as session:
        assert cleanup_stale_verification_runs(session, storage)["runs_cleaned"] == 1
        assert {str(row.id): row.evidence_sha256 for row in session.scalars(select(VerificationReportRecord))} == saved


def test_transient_retry_respects_persisted_task_budget(monkeypatch):
    from celery.exceptions import Retry
    from datetime import timedelta
    from app.services import paper_workflow
    from app.services.paper_dispatch import DISPATCH_KEY, attempt_id_for
    from app.tasks import check_paper as tasks
    factory, storage, job_id = _retry_test_job(monkeypatch)
    monkeypatch.setattr(tasks, "SessionLocal", factory)
    monkeypatch.setattr(tasks, "get_storage_backend", lambda: storage)
    monkeypatch.setattr(paper_workflow.settings, "PAPER_LLM_PROCESSING_ENABLED", False)

    def fail(*args, **kwargs):
        raise ConnectionError("synthetic temporary processing failure")

    def retry(**kwargs):
        raise Retry()

    monkeypatch.setattr(paper_workflow, "_shadow_artifact", fail)
    task = SimpleNamespace(request=SimpleNamespace(retries=0), retry=retry)
    with factory() as session:
        attempt = attempt_id_for(session.get(Job, job_id))
        assert attempt
    for count in range(1, 5):
        with pytest.raises(Retry if count < 4 else RuntimeError):
            tasks._run_stage(task, str(job_id), attempt, "verify")
        with factory() as session:
            job = session.get(Job, job_id)
            intent = job.upload_evidence[DISPATCH_KEY]
            assert intent["stage_starts"]["verify"] == count
            assert job.verification_summary is None
            run = session.scalar(select(VerificationRunRecord))
            assert run.status == ("active" if count < 4 else "failed_cleaned")
            if count < 4:
                due = (datetime.now(timezone.utc) - timedelta(seconds=1)).isoformat()
                job.upload_evidence = {**job.upload_evidence, DISPATCH_KEY: {
                    **intent, "stage_retry_after": {"verify": due}, "dispatch_after": due,
                }}
                session.commit()
            else:
                assert job.status == "failed"
    assert storage.list_keys("verification-runs/") == []


def test_transient_retry_abandoned_bytes_expire(monkeypatch):
    from datetime import timedelta
    from app.services import paper_workflow
    from app.services.verification_run import cleanup_stale_verification_runs
    factory, storage, job_id = _retry_test_job(monkeypatch)

    def fail(*args, **kwargs):
        raise ConnectionError("synthetic abandoned retry")

    monkeypatch.setattr(paper_workflow, "_shadow_artifact", fail)
    with pytest.raises(ConnectionError):
        verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        assert cleanup_stale_verification_runs(session, storage)["runs_cleaned"] == 0
        run = session.scalar(select(VerificationRunRecord))
        assert cleanup_stale_verification_runs(session, storage,
            now=run.lease_expires_at + timedelta(seconds=1))["runs_cleaned"] == 1
        assert run.status == "abandoned_cleaned"
    assert storage.list_keys("verification-runs/") == []


def test_verification_retry_recovers_committed_packages_after_cleanup(monkeypatch):
    from app.services import paper_workflow
    factory, storage, job_id = _retry_test_job(monkeypatch)
    original = paper_workflow._citation_group_index
    def fail_once(*args, **kwargs):
        monkeypatch.setattr(paper_workflow, "_citation_group_index", original)
        raise ConnectionError("failure after source completion")
    monkeypatch.setattr(paper_workflow, "_citation_group_index", fail_once)
    with pytest.raises(ConnectionError):
        verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        saved = {str(r.id): r.evidence_sha256 for r in session.scalars(select(VerificationReportRecord))}
        assert len(saved) == 2
        assert session.scalar(select(VerificationRunRecord)).status == "cleaned"
    monkeypatch.setattr(paper_workflow, "_shadow_artifact", lambda *a, **kw: pytest.fail("saved evidence must not rerun"))
    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert summary["reports_persisted"] == 2
    assert summary["source_failures"] == []
    assert set(summary["report_ids"]) == set(saved)
    with factory() as session:
        assert {str(r.id): r.evidence_sha256 for r in session.scalars(select(VerificationReportRecord))} == saved


@pytest.mark.parametrize("damage", ["missing", "payload", "scope", "source", "claim"])
def test_verification_recovery_refuses_damaged_or_retargeted_evidence(monkeypatch, damage):
    from app.services.paper_workflow import PaperWorkflowError
    from app.services.verification_report import _payload_digest
    from app.services.evidence_package import _package_payload_sha256
    factory, storage, job_id = _retry_test_job(monkeypatch)
    verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        job = session.get(Job, job_id)
        job.stage = "verifying"
        job.verification_summary = None  # simulate uncommitted aggregate checkpoint
        record = session.scalars(select(VerificationReportRecord)).first()
        if damage == "missing":
            session.delete(record)
        elif damage == "scope":
            record.scope_id = "another-owner"
        else:
            payload = json.loads(json.dumps(record.report_payload))
            if damage == "payload":
                payload["verdict"] = "tampered"
            else:
                package = payload["authoritative_evidence_package"]
                if damage == "source":
                    package["source_identity"]["content_sha256"] = "0" * 64
                else:
                    package["student_text"] = "A different student's wording."
                package["package_sha256"] = _package_payload_sha256({k: v for k, v in package.items() if k != "package_sha256"})
                record.evidence_sha256 = _payload_digest(payload)
            record.report_payload = payload
        session.commit()
    with pytest.raises(PaperWorkflowError) as caught:
        verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert caught.value.code in {"verification_report_binding_invalid", "verification_report_missing"}


def test_cleanup_pending_does_not_hide_successfully_saved_evidence(monkeypatch):
    from app.services.verification_run import cleanup_stale_verification_runs
    factory, storage, job_id = _retry_test_job(monkeypatch)
    original = storage.delete
    monkeypatch.setattr(storage, "delete", lambda key: False if key.startswith("verification-runs/") else original(key))
    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert summary["reports_persisted"] == 2 and summary["source_failures"] == []
    with factory() as session:
        assert session.scalar(select(VerificationRunRecord)).status == "cleanup_pending"
    monkeypatch.setattr(storage, "delete", original)
    with factory() as session:
        assert cleanup_stale_verification_runs(session, storage)["runs_cleaned"] == 1
        assert session.scalar(select(VerificationRunRecord)).status == "cleaned"


def test_finalization_refuses_a_missing_immutable_package(monkeypatch):
    from app.services.paper_workflow import PaperWorkflowError
    factory, storage, job_id = _retry_test_job(monkeypatch)
    verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        session.delete(session.scalars(select(VerificationReportRecord)).first())
        session.commit()
    with pytest.raises(PaperWorkflowError), factory() as session:
        finalize_paper_job(session, storage, job_id)
    with factory() as session:
        assert session.scalars(select(Report)).all() == []


def test_collective_claim_is_verified_against_each_exact_source_membership():
    claim = ClaimEvidence(
        claim_id="collective-claim",
        paper_version_id="paper-v1",
        text="Smith and Jones show that the result is stable.",
        reference_ids=["ref-smith", "ref-jones"],
        passage_start=0,
        passage_end=47,
    )

    grouped = _claims_by_reference([claim])

    assert grouped == {"ref-smith": [claim], "ref-jones": [claim]}


def test_unbound_continuation_preserves_anchor_without_blocking_explicit_evidence(monkeypatch):
    from app.services import paper_workflow
    factory, storage, job_id = _retry_test_job(monkeypatch)
    with factory() as session:
        job = session.get(Job, job_id)
        payload = dict(job.extraction_payload)
        claims = [dict(item) for item in payload["citation_claims"]]
        assert len(claims) == 2
        claims[1].update(citation_marker="implicit_continuation", citation_markers=[])
        continuation_id = claims[1]["claim_id"]
        payload["citation_claims"] = claims
        job.extraction_payload = payload
        session.commit()
    checked = []
    original = paper_workflow._shadow_artifact
    def capture(source, claim, *args, **kwargs):
        checked.append(claim.claim_id)
        return original(source, claim, *args, **kwargs)
    monkeypatch.setattr(paper_workflow, "_shadow_artifact", capture)
    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert checked == [claims[0]["claim_id"]]
    assert summary["reports_persisted"] == 1
    assert not summary["source_failures"]
    continuation = next(g for g in summary["citation_groups"] if g["claim_id"] == continuation_id)
    assert continuation["members"][0]["status"] == "citation_not_assessed"
    assert continuation["members"][0]["reason_code"] == "continuation_source_marker_unresolved"
    with factory() as session:
        finalize_paper_job(session, storage, job_id)
        view = session.scalar(select(Report)).report_json["evidence_report"]
    citation = next(c for c in view["citations"] if c["claim_id"] == continuation_id)
    assert citation["student_text"] == claims[1]["text"]
    member = citation["members"][0]
    assert member["status"] == "citation_not_assessed"
    assert not member["best_evidence"]
    assert not member["show_quotation_check"]
    html = _render_member(member)
    assert "Continuation not assessed" in html
    assert "Source not retrieved" not in html
    panel = _render_panel_template(citation, 1)
    assert '<h3 class="member-label">Not assessed</h3>' in panel
    assert "<h3>Not retrieved</h3>" not in panel


def test_report_distinguishes_failed_cited_webpage_with_metadata_only():
    assert _report_source_failure_reason(
        {
            "reason_code": "source_not_found",
            "reference_discovery_trace": {
                "attempts": [
                    {
                        "route_category": "student_url",
                        "outcome": "operational_failure",
                    },
                    {
                        "route_category": "academic_adapter",
                        "outcome": "candidate_found",
                    },
                ]
            },
        }
    ) == "cited_webpage_unavailable_metadata_only"


def test_only_blocking_providers_become_retry_dependencies():
    """A provider is a re-run trigger only if the completion gate counts it.

    Rewritten 2026-09-24. The previous version asserted that SearXNG became a
    trigger because its query was cooldown-skipped, whatever its role. All four
    re-run waves on 2026-09-23 -- 192 reference re-runs -- came from exactly
    that, though SearXNG cannot complete or invalidate a search under
    `api-first-search-v2`, and they moved fabrication flags between report
    versions without changing any completion gate.
    """
    from tests.unit.blocked_discovery import blocked_record

    optional = blocked_record("r1", policy="api-first-search-v2", web_outcomes={
        "brave": "results", "exa": "results", "searxng": "cooldown_skipped"})
    required_web = blocked_record("r2", policy="api-first-search-v2",
                                  web_outcomes={"brave": "timeout", "exa": "results"})
    adapter = blocked_record("r3", adapter="core")

    assert _retryable_provider_dependencies(None, optional) == []
    assert _retryable_provider_dependencies(None, required_web) == ["brave"]
    assert _retryable_provider_dependencies(None, adapter) == ["core"]
    # A reference that got a real answer is never re-run.
    assert _retryable_provider_dependencies(
        None, {**adapter, "outcome": "unlocated_after_search"}) == []
    # An unreadable record names nothing rather than guessing.
    assert _retryable_provider_dependencies(None, {"outcome": "search_incomplete"}) == []


@pytest.mark.parametrize("identity_found", [True, False])
def test_metadata_only_identity_survives_checkpoint_reload_and_report(monkeypatch, identity_found):
    from app.services.reference_discovery import ExpectedBibliographicFields
    from app.services.source_resolver import SourceResolver, SourceResolutionError, _ACTIVE_DISCOVERY_TRACE

    def forbidden(*args, **kwargs):
        pytest.fail("Network is forbidden in metadata-only workflow regression")

    monkeypatch.setattr("socket.socket.connect", forbidden)
    monkeypatch.setattr("socket.getaddrinfo", forbidden)
    monkeypatch.setattr(paper_upload, "scan_with_clamd", lambda _: (SafetyVerdict.CLEAN, "OK"))
    engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    captured = {}

    class MetadataOnlyResolver:
        calls = 0

        def resolve_reference(self, reference):
            self.calls += 1
            token = _ACTIVE_DISCOVERY_TRACE.set({
                "reference_id": reference.reference_id,
                "expected": ExpectedBibliographicFields(title=reference.title,
                    authors=[reference.author], year=reference.year),
                "search_policy_version": "api-first-search-v2",
                "required_web_providers": ["brave", "exa"],
                "required": {"academic_adapter", "bounded_web"},
                "queries": [], "attempts": [], "candidates": [], "limitations": [],
            })
            try:
                # Independently declared synthetic metadata; never source text.
                SourceResolver._record_discovery_attempt(category="academic_adapter", provider="crossref",
                    result=RetrievalResult(source_name="crossref", success=identity_found,
                        title="A useful title" if identity_found else None,
                        authors=["Smith, J."] if identity_found else [],
                        year="2020" if identity_found else None, error=None if identity_found else "No results"),
                    required=True)
                trace, discovery = SourceResolver._discovery_artifacts()
            finally:
                _ACTIVE_DISCOVERY_TRACE.reset(token)
            captured.update(trace=trace, discovery=discovery)
            exc = SourceResolutionError("Synthetic acquisition unavailable")
            exc.reference_discovery_trace = trace
            exc.reference_discovery = discovery
            raise exc

    resolver = MetadataOnlyResolver()
    with factory() as session:
        job = create_paper_job(session, storage, content=_paper_bytes(), filename="identity.docx",
            media_type=DOCX_MEDIA_TYPE, scope_id="personal-default")
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        assert retrieve_paper_sources(session, storage, job_id, resolver=resolver) == {"unavailable": 1}
    expected = "confirmed" if identity_found else "search_incomplete"
    with factory() as session:
        job = session.get(Job, job_id)
        result = job.source_results[0]
        assert result["reference_discovery"] == captured["discovery"]
        assert result["reference_discovery_trace"] == captured["trace"]
        assert result["reference_discovery"]["outcome"] == expected
        assert result["status"] == "unavailable"
        assert not result.get("representation_id") and not result.get("abstract_evidence")
        retrieve_paper_sources(session, storage, job_id, resolver=resolver)
        assert resolver.calls == 1
    verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        finalize_paper_job(session, storage, job_id)
    with factory() as session:
        job = session.get(Job, job_id)
        report = session.scalar(select(Report).where(Report.job_id == job_id))
        view = report.report_json["evidence_report"]
        assert get_authorized_evidence_report_view(session, report.id,
            scope_type="personal_owner", scope_id="personal-default") == view
        with pytest.raises(EvidenceReportAuthorizationError):
            get_authorized_evidence_report_view(session, report.id,
                scope_type="personal_owner", scope_id="another-owner")
        assert view["overview"]["verified_full_text_sources"] == 0
        assert session.scalars(select(VerificationReportRecord)).all() == []
        assert len(view["citations"]) == 2
        for citation in view["citations"]:
            assert len(citation["members"]) == 1
            for member in citation["members"]:
                assert member["reference_identity"]["status"] == expected
                assert not member["reference_identity"]["attention"]
                assert not member.get("best_evidence")
        assert job.source_results[0]["reference_discovery_trace"] == captured["trace"]
    engine.dispose()


@pytest.mark.parametrize("failure", [False, True])
@pytest.mark.parametrize("observed", [False, True])
def test_submitted_link_observations_checkpoint_and_report_data(monkeypatch, failure, observed):
    import httpx
    from app.services.submitted_links import observe_reference, observe_request, response_observed
    from app.services.source_resolver import SourceResolutionError
    monkeypatch.setattr(paper_upload, "scan_with_clamd", lambda _: (SafetyVerdict.CLEAN, "OK"))
    engine = create_engine("sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    document = Document()
    document.add_paragraph("Careful verification improves accuracy (Smith, 2020).")
    document.add_paragraph("References")
    document.add_paragraph("Smith, J. (2020). A useful title. Example Press. https://example.org/source")
    data = io.BytesIO(); document.save(data)
    class ObservedResolver(Resolver):
        @observe_reference
        def resolve_reference(self, reference):
            if observed:
                @observe_request
                def fixture_request(url):
                    response_observed(url, 404)
                    return httpx.Response(404, content=b"fixture only")
                fixture_request(reference.url)
            if failure:
                self.calls += 1
                raise SourceResolutionError("fixture unavailable")
            return super().resolve_reference(reference)
    resolver = ObservedResolver()
    with factory() as session:
        job = create_paper_job(session, storage, content=data.getvalue(), filename="fixture.docx",
            media_type=DOCX_MEDIA_TYPE, scope_id="personal-default")
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        retrieve_paper_sources(session, storage, job_id, resolver=resolver)
        expected = job.source_results[0]["submitted_link_observations"]
        assert expected and all(r["state"] == ("observed" if observed else "not_checked") for r in expected)
        if observed:
            assert expected[0]["requests"][0]["outcome"] == "not_found"
    with factory() as session:
        job = session.get(Job, job_id)
        assert job.source_results[0]["submitted_link_observations"] == expected
        retrieve_paper_sources(session, storage, job_id, resolver=resolver)
        assert resolver.calls == 1
    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert summary["submitted_link_observations"] == expected
    with factory() as session:
        finalize_paper_job(session, storage, job_id)
    with factory() as session:
        report = session.scalar(select(Report).where(Report.job_id == job_id))
        assert report.report_json["submitted_link_observations"] == expected
        view = report.report_json['evidence_report']
        assert view['submitted_link_observations'] == expected
        member_rows = [row for citation in view['citations'] for member in citation['members']
                       for row in member.get('submitted_link_observations', [])]
        assert member_rows == expected
        from app.services.evidence_report import render_evidence_report_html
        for audience in ('student', 'instructor'):
            html = render_evidence_report_html({**view, 'audience': audience}, csp_nonce='link-check-test-nonce')
            assert 'Submitted links in uncited references' not in html
            assert 'Submitted link check' not in html  # No localized reference flag in this raw view.
    engine.dispose()


@pytest.mark.parametrize('require_links',[False, True])
def test_checkpointed_workflow_persists_shadow_report_and_cleans(monkeypatch, require_links):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_paper_bytes(),
            filename="paper.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
            require_reference_links=require_links,
        )
        job_id = job.id
        extracted = extract_paper_job(session, storage, job_id, llm_enabled=False)
        assert extracted["citation_units"] == 2
        assert extracted["reference_layout_status"] == "complete"
        assert extracted["reference_layout_matched"] == extracted["reference_layout_total"]
        assert extracted["reference_formatting_status"] == "partial"
        assert extracted["reference_formatting_differences"] == 1
        job.upload_evidence = {**job.upload_evidence,
            'assessment_configuration':{'version':'assessment-configuration-v1', 'require_reference_links':not require_links}}
        session.commit()
        assert job.input_storage_key is None
        assert extract_paper_job(session, storage, job_id, llm_enabled=False) == extracted
        class PrefetchResolver(Resolver):
            def __init__(self):
                super().__init__()
                self.phases = []

            def prefetch_title_candidates(self, queries):
                self.phases.append('titles')
                assert queries
                return {}

            def prefetch_deferred_dois(self, dois):
                assert self.phases == ['titles']
                self.phases.append('dois')
                return {}

            def resolve_reference(self, reference):
                assert self.phases == ['titles', 'dois']
                self.phases.append('resolve')
                return super().resolve_reference(reference)

        resolver = PrefetchResolver()
        retrieved = retrieve_paper_sources(
            session,
            storage,
            job_id,
            resolver=resolver,
        )
        assert retrieved == {"transient_authorized": 1}
        assert job.source_results[0]["reference_discovery_trace"] == {
            "trace_version": "reference-discovery-trace-v1",
            "outcome_derived": True,
        }
        assert job.source_results[0]["reference_discovery"] == {
            "record_version": "reference-discovery-v1",
            "outcome": "confirmed",
        }
        assert resolver.calls == 1
        assert retrieve_paper_sources(session, storage, job_id, resolver=resolver) == retrieved
        assert resolver.calls == 1

    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert summary["reports_persisted"] == 2
    assert len(summary["citation_groups"]) == 2
    assert all(
        group["expected_member_count"] == 1
        and group["persisted_member_count"] == 1
        and group["coverage_status"] == "complete"
        for group in summary["citation_groups"]
    )
    assert summary["relationship_mode"] == "shadow"
    assert summary["decision_applied"] is False
    assert verify_paper_sources(factory, storage, job_id, llm_enabled=False) == summary

    with factory() as session:
        finalized = finalize_paper_job(session, storage, job_id)
        job = session.get(Job, job_id)
        evidence = session.scalars(select(VerificationReportRecord)).all()
        aggregate = session.scalar(select(Report))
        run = session.scalar(select(VerificationRunRecord))
        assert finalized["input_cleaned"] is True
        assert job.status == "completed"
        assert job.stage == "completed"
        assert job.input_storage_key is None
        assert len(evidence) == 2
        assert all(item.verdict == "inconclusive" for item in evidence)
        assert all(
            item.report_payload["verification_candidates"]["status"] == "complete"
            for item in evidence
        )
        assert all(
            item.report_payload["candidate_relationships"]["status"] == "not_run"
            for item in evidence
        )
        assert all(
            item.report_payload["facet_evidence_foundation"]["status"] == "complete"
            for item in evidence
        )
        assert all(
            item.report_payload["facet_evidence_ledger"]["status"] == "not_run"
            for item in evidence
        )
        assert aggregate.report_json["decision_applied"] is False
        assert aggregate.report_json["reference_consistency"]["assessment_version"] == "reference-consistency-v3"
        assert aggregate.report_json["reference_consistency"]["formatting_status"] == "not_assessed"
        assert aggregate.report_json["reference_consistency"]["formatting_reason_codes"] == [
            "layout_evidence_ready_style_rules_not_accepted"
        ]
        assert aggregate.report_json["reference_layout"]["status"] == "complete"
        assert aggregate.report_json["reference_formatting"]["assessment_version"] == (
            "reference-formatting-v2"
        )
        assert aggregate.report_json["reference_formatting"]["result_counts"] == {
            "difference": 1
        }
        evidence_view = aggregate.report_json["evidence_report"]
        assert evidence_view['assessment_configuration']['require_reference_links'] is require_links
        assert evidence_view["submitted_locator_inventory"] is not None
        assert evidence_view["submitted_locator_inventory"]["contributes_to_issue_counts"] is False
        assert evidence_view["view_version"] == "evidence-led-report-v31"
        assert evidence_view['reference_credibility']
        assert all(not item['findings'] for item in evidence_view['reference_credibility'].values())
        assert evidence_view["paper_surface"]["status"] == "presentation_source_retained"
        assert evidence_view["paper_surface"]["reason_code"] == (
            "deterministic_pdf_render_required"
        )
        artifact = session.scalar(select(ReportPaperArtifactRecord))
        assert artifact.report_id == aggregate.id
        loaded_record, loaded_bytes = load_authorized_report_paper_artifact(
            session,
            storage,
            report_id=aggregate.id,
            artifact_id=artifact.id,
            scope_type="personal_owner",
            scope_id="personal-default",
        )
        assert loaded_record.id == artifact.id
        assert loaded_bytes != _paper_bytes()
        sanitized_document = Document(io.BytesIO(loaded_bytes))
        assert sanitized_document.paragraphs[0].text == (
            "Careful verification improves accuracy (Smith, 2020)."
        )
        with pytest.raises(ReportPaperArtifactError):
            load_authorized_report_paper_artifact(
                session,
                storage,
                report_id=aggregate.id,
                artifact_id=artifact.id,
                scope_type="personal_owner",
                scope_id="another-owner",
            )
        assert evidence_view["overview"]["citations_analyzed"] == 2
        assert evidence_view["overview"]["verified_full_text_sources"] == 1
        assert evidence_view["overview"]["limited_or_unavailable_sources"] == 0
        assert all(
            citation["members"][0]["reference_identity"]["status"] == "confirmed"
            for citation in evidence_view["citations"]
        )
        assert evidence_view["citation_use_assessment"] is None
        serialized_view = json.dumps(evidence_view)
        assert "candidate_relationships" not in serialized_view
        assert "structured_judgment" not in serialized_view
        assert "retrieval_score" not in serialized_view
        rendered = render_evidence_report_html(
            evidence_view, csp_nonce="test-report-nonce-1234"
        )
        assert "Evidence-led report" not in rendered
        assert "paper.docx" in rendered
        assert "Careful verification improves accuracy" in rendered
        assert "structured_judgment" not in rendered
        assert "Candidate-specific retrieval" not in rendered
        assert (
            get_authorized_evidence_report_view(
                session,
                aggregate.id,
                scope_type="personal_owner",
                scope_id="personal-default",
            )
            == evidence_view
        )
        with pytest.raises(EvidenceReportAuthorizationError):
            get_authorized_evidence_report_view(
                session,
                aggregate.id,
                scope_type="personal_owner",
                scope_id="another-owner",
            )
        assert run.status == "cleaned"
        assert run.edition_or_version == "author_accepted_manuscript"
        assert run.transient_objects == []
        assert finalize_paper_job(session, storage, job_id)["report_id"] == finalized["report_id"]
    _assert_only_report_marking_copy(storage)


def test_evidence_report_renderer_escapes_untrusted_student_and_source_text():
    view = {
        "title": '<script id="title-attack">alert(1)</script>',
        "citation_format": "APA",
        "paper_surface": {"message": '<img src=x onerror="alert(2)">'},
        "overview": {
            "citations_analyzed": 1,
            "verified_full_text_sources": 1,
            "limited_or_unavailable_sources": 0,
            "reference_identity_attention": 0,
            "quotation_differences_attention": 0,
            "locator_differences_attention": 0,
        },
        "citations": [
            {
                "tone": "evidence_available",
                "student_text": '<svg onload="alert(3)">',
                "members": [
                    {
                        "source": {
                            "author": "Smith",
                            "year": "2020",
                            "title": '<iframe src="bad"></iframe>',
                            "raw_reference": "Smith. Journal Name. <script>bad</script>",
                            "text_style_spans": [
                                {"start": 7, "end": 19, "italic": True, "bold": False}
                            ],
                        },
                        "availability": "Verified source representation available",
                        "best_evidence": {
                            "text": '<script id="source-attack">alert(4)</script>',
                            "locator": "Page 1",
                        },
                        "additional_evidence": [],
                        "quotation_check": {"attention": False, "label": "Not applicable"},
                        "locator_check": {"attention": False, "label": "Not applicable"},
                        "limitations": [],
                    }
                ],
            }
        ],
        "limits": [],
    }

    rendered = render_evidence_report_html(
        view, csp_nonce="test-report-nonce-1234"
    )

    assert '<script id="title-attack">' not in rendered
    assert '<script id="source-attack">' not in rendered
    assert '<svg onload="alert(3)">' not in rendered
    assert '<iframe src="bad">' not in rendered
    assert "&lt;script id=&quot;title-attack&quot;&gt;" in rendered
    assert "<em>Journal Name</em>" in rendered
    assert "&lt;script&gt;bad&lt;/script&gt;" in rendered


def test_evidence_report_renders_continuous_page_surface_with_typed_overlay():
    view = {
        "title": "Page-first report",
        "citation_format": "APA",
        "paper_surface": {
            "message": "Page-faithful paper retained.",
            "page_dimensions": [
                {"page_index": 0, "width": 612.0, "height": 792.0}
            ],
            "page_href_template": "/report/report-1/paper/page/{page_index}",
        },
        "overview": {
            "citations_analyzed": 1,
            "verified_full_text_sources": 0,
            "limited_or_unavailable_sources": 1,
            "reference_identity_attention": 0,
            "quotation_differences_attention": 0,
            "locator_differences_attention": 0,
        },
        "citations": [
            {
                "claim_id": "claim-1",
                "tone": "not_assessed",
                "student_text": "A citation (Smith, 2020).",
                "members": [],
                "paper_location": {
                    "localization_level": "exact_rectangle",
                    "rectangles": [
                        {
                            "page_index": 0,
                            "x0": 72.0,
                            "y0": 100.0,
                            "x1": 240.0,
                            "y1": 118.0,
                        }
                    ],
                    "action": None,
                },
            }
        ],
        "limits": [],
    }

    rendered = render_evidence_report_html(
        view, csp_nonce="page-first-report-nonce"
    )

    assert 'href="/report/report-1/paper/page/0"' in rendered
    assert 'class="citation-overlay not_assessed coverage_none"' in rendered
    assert 'class="selection-bg"' in rendered
    assert 'class="underline"' in rendered
    assert "stroke-dasharray:2 3" not in rendered
    assert 'class="selected-citation"' in rendered
    assert "View in paper" not in rendered
    assert "Citation evidence</h2>" not in rendered
    assert "Submitted paper</h2>" not in rendered
    assert 'href="#citation-location-1">Citation 1</a><span data-proposition-suffix></span></h2>' in rendered
    assert "Citations without exact page geometry" not in rendered
    assert "mouseenter" not in rendered
    assert 'id="paper-layout"' not in rendered and 'id="sources-layout"' not in rendered
    assert 'id="sources-layout"' not in rendered   # one layout (2026-09-28)
    assert 'id="judgment-layout"' not in rendered
    assert "[hidden]{display:none!important}" in rendered
    assert 'id="report-splitter"' in rendered
    assert 'aria-valuemin="30" aria-valuemax="80" aria-valuenow="68"' in rendered
    assert 'aria-label="Resize paper and evidence panels"' in rendered
    assert "Paper view" not in rendered
    assert "Citation spans represented" not in rendered
    assert "How to read this report" in rendered
    assert "Patterns and issues" not in rendered and "Priorities for revision" not in rendered
    assert "Repeated patterns" not in rendered
    assert "Report Counts and Evidence Breakdown" not in rendered
    assert "Facet Fidelity" not in rendered
    assert 'id="zoom-in"' in rendered and 'id="zoom-out"' in rendered
    assert 'id="add-comment"' not in rendered
    assert 'id="add-highlight"' not in rendered
    assert 'id="pen-tool"' not in rendered
    assert 'id="undo-annotation"' not in rendered
    assert 'id="redo-annotation"' not in rendered
    assert "solid underline" not in rendered
    assert "dashed: abstract or limited text retrieval" not in rendered
    assert "dotted: no text retrieved" not in rendered
    assert 'class="color-key teal"></i>abstract' not in rendered   # swatches removed 2026-09-28
    assert 'class="color-key mauve"></i>limited text' not in rendered
    assert "brown: checked source; no clear matching passage" not in rendered
    assert 'class="color-key gray"></i>no retrieved text' not in rendered
    assert "Light-grey underline: complete citation span" not in rendered   # key rebuilt 2026-09-28
    assert ".citation-overlay.retrieved_no_connection{" in rendered
    assert ".citation-overlay.retrieved_no_connection .underline" not in rendered


def test_stored_annotations_are_never_rendered():
    """The comment, highlight and pen tools were removed (owner decision
    2026-09-25); records stored before then stay untouched and unshown."""
    anchor_id = "a" * 64
    view = {
        "title": "Annotated report",
        "citation_format": "APA",
        "paper_surface": {
            "message": "Page-faithful paper retained.",
            "page_dimensions": [{"page_index": 0, "width": 612.0, "height": 792.0}],
            "page_href_template": "/report/report-1/paper/page/{page_index}",
        },
        "overview": {},
        "citations": [
            {
                "claim_id": "claim-1",
                "tone": "evidence_available",
                "display_coverage": "full",
                "student_text": "A citation (Smith, 2020).",
                "members": [],
                "paper_location": {
                    "anchor_id": anchor_id,
                    "localization_level": "exact_rectangle",
                    "rectangles": [
                        {
                            "page_index": 0,
                            "x0": 72.0,
                            "y0": 100.0,
                            "x1": 240.0,
                            "y1": 118.0,
                        }
                    ],
                },
            }
        ],
        "annotations": [
            {
                "annotation_id": "b" * 36,
                "revision": 2,
                "annotation_type": "comment",
                "anchor": {"anchor_id": anchor_id},
                "content": "Check this distinction.",
                "visibility": "released",
            },
            {
                "annotation_id": "c" * 36,
                "revision": 1,
                "annotation_type": "highlight",
                "anchor": {"anchor_id": anchor_id},
                "content": None,
                "visibility": "private",
            },
        ],
        "annotation_action": {
            "create_href": "/report/report-1/annotations",
            "revision_href_template": "/report/report-1/annotations/{annotation_id}",
        },
        "limits": [],
    }

    rendered = render_evidence_report_html(view, csp_nonce="saved-annotation-nonce")

    for absent in ("persisted-highlight", "instructor-comment-marker", "Check this distinction.",
                   "data-annotation-operation", "instructor comment or highlight", "annotation-panel-",
                   "data-annotation-create", "Pen (draft)"):
        assert absent not in rendered


def test_abstract_only_member_shows_abstract_once_without_unavailable_or_checks():
    reference = SimpleNamespace(
        reference_id="ref-abstract",
        author="Dearden, J.",
        year="2014",
        title="English as a medium of instruction",
        raw_ref="Dearden, J. (2014). English as a medium of instruction.",
        doi="",
        url="",
    )
    claim = SimpleNamespace(claim_type="quotation", page_locator="12")
    member = _unavailable_member(
        {
            "abstract_available": True,
            "abstract_evidence": {
                "text": "The abstract describes the study and its main scope.",
                "truncated": False,
            },
            "reason_code": "full_text_unavailable",
        },
        reference,
        claim,
    )

    rendered = _render_member(member)

    # Collapsed under "Abstract" (owner request 2026-09-28).
    assert ('<details class="abstract-disclosure"><summary>Abstract</summary><blockquote>'
            'The abstract describes the study') in rendered
    assert "Abstract evidence" not in rendered
    assert "Only the abstract was retrieved" not in rendered
    assert 'class="source-excerpt"' not in rendered
    assert 'class="locator">Abstract' not in rendered
    assert "No source evidence is available" not in rendered
    assert "Deterministic checks" not in rendered
    assert "Quotation and locator information" not in rendered
    assert member["limitations"] == []


def test_partial_selection_label_is_not_projected_as_support_judgment():
    passage = {
        "passage_id": "partial",
        "excerpt": "Participants reported that general training opportunities were limited.",
        "page_label": "4",
    }

    view = _passage_view(
        passage,
        {
            "passage_id": "partial",
            "relevance": "partially_relevant",
            "confidence": "medium",
            "evidence_role": "source_own_claim_or_finding",
        },
    )

    assert "addresses only part of the citation" not in view["evidence_note"]
    assert "does not establish every material detail" not in view["evidence_note"]


def test_minor_quotation_difference_names_and_marks_the_changed_word():
    claim = (
        'Using thematic analysis, a method “for identifying, analysing, and '
        'reporting patterns (themes) in data” (Braun & Clarke, 2006, p. 79).'
    )
    source = (
        "Thematic analysis is a method for identifying, analysing, and "
        "reporting patterns (themes) within data."
    )

    differences = _quotation_difference_diagnostics(claim, [source])

    assert len(differences) == 1
    assert differences[0]["severity"] == "minor"
    assert differences[0]["changed_word_count"] == 1
    assert differences[0]["spans"][0]["paper_text"] == "in"
    assert differences[0]["spans"][0]["source_text"] == "within"
    assert "1 of 9 words" in _quotation_difference_label(differences)


def test_display_excerpt_trims_byline_and_leads_with_responsive_sentence():
    value = (
        "An Article Title. Author Name, Example University. Email: author@example.test. "
        "Abstract This study provides general background. The participants reported "
        "that targeted language training was unavailable to their instructors. "
        "The paper then discusses an unrelated administrative issue in detail. "
        "Further unrelated context follows for several paragraphs. "
    ) * 4

    displayed = _responsive_display_excerpt(
        value,
        "Instructors often do not receive targeted language training.",
    )

    assert "An Article Title" not in displayed
    assert "targeted language training" in displayed
    assert len(displayed) <= 900


def test_display_excerpt_drops_unresponsive_neighbor_and_speaker_label():
    value = (
        "Unrelated students discussed their general English proficiency. "
        "(Instructor, partner institution) Participants reported that targeted "
        "language training was unavailable to their instructors. "
        "A separate administrative concern followed this finding. "
    ) * 5

    displayed = _responsive_display_excerpt(
        value,
        "Instructors often do not receive targeted language training.",
    )

    assert displayed.startswith("Participants reported")
    assert "general English proficiency" not in displayed


def test_primary_evidence_uses_one_or_two_sentences_and_retains_full_context():
    passage = {
        "passage_id": "responsive",
        "excerpt": (
            "The introduction summarizes earlier work on language policy. "
            "Participants reported that targeted language training was unavailable to instructors. "
            "The following section describes the survey software. "
            "A final sentence discusses administrative scheduling."
        ),
        "page_label": "8",
    }

    view = _passage_view(
        passage,
        {
            "passage_id": "responsive",
            "relevance": "relevant",
            "confidence": "high",
            "evidence_role": "source_own_claim_or_finding",
        },
        claim_text="Instructors did not receive targeted language training.",
    )

    assert view["display_text"].startswith("Participants reported")
    assert "survey software" not in view["display_text"]
    assert "survey software" in view["context_text"]


def test_compound_panel_has_no_citation_information_section():
    def member(author):
        return {
            "source": {
                "author": author,
                "year": "2024",
                "title": "A source",
                "raw_reference": f"{author}. (2024). A source.",
                "text_style_spans": [],
            },
            "availability": "",
            "best_evidence": None,
            "additional_evidence": [],
            "reference_identity": {"attention": False},
            "show_quotation_check": False,
            "show_locator_check": False,
            "limitations": ["Source completeness is uncertain. Check the source manually."],
        }

    rendered = _render_panel_template(
        {
            "citation_number": 7,
            "student_text": "A claim with two cited sources.",
            "members": [member("First"), member("Second")],
            "boundary_reason": "Source membership is resolved.",
        },
        7,
    )

    # Removed from the window (owner request 2026-09-28).
    assert "Citation Information" not in rendered
    assert "Source completeness is uncertain" not in rendered
    assert "Citation 7" in rendered
    assert "Material Limitation" not in rendered


def test_insufficiency_display_priority_keeps_direct_student_self_assessment_first():
    passages = [
        {"excerpt": "Lecturers reported that their own English was inadequate."},
        {
            "excerpt": (
                "The collective picture is one of deep concern about students' English level. "
                "This is matched by student evaluations of their own English proficiency."
            )
        },
    ]

    ordered = _prioritize_display_passages(
        passages,
        "Most students feel their English level is not sufficient.",
    )

    assert ordered[0] is passages[1]


def test_display_priority_uses_persisted_relevance_without_changing_union():
    passages = [
        {
            "passage_id": "topic",
            "excerpt": "The article surveys English-medium instruction.",
            "retrieval_score": 0.95,
            "boundary_status": "sentence_complete",
        },
        {
            "passage_id": "direct",
            "excerpt": "Students reported that their English proficiency was inadequate.",
            "retrieval_score": 0.61,
            "boundary_status": "sentence_complete",
        },
    ]
    gate = {
        "assessments": [
            {
                "passage_id": "topic",
                "relevance": "partially_relevant",
                "confidence": "high",
            },
            {
                "passage_id": "direct",
                "relevance": "relevant",
                "confidence": "high",
            },
        ]
    }

    ordered = _prioritize_display_passages(
        passages,
        "Students feel their English level is insufficient.",
        gate,
    )

    assert [item["passage_id"] for item in ordered] == ["direct", "topic"]
    assert {item["passage_id"] for item in ordered} == {"direct", "topic"}


def test_indirect_passage_reaches_additive_selection_when_direct_evidence_exists():
    passages = [
        {"passage_id": "direct", "excerpt": "The source reports its own finding."},
        {"passage_id": "indirect", "excerpt": "Another study reports the idea."},
    ]
    gate = {
        "status": "complete",
        "assessments": [
            {
                "passage_id": "direct",
                "relevance": "relevant",
                "evidence_role": "source_own_claim_or_finding",
            },
            {
                "passage_id": "indirect",
                "relevance": "relevant",
                "evidence_role": "representation_of_other_work",
            },
        ],
    }

    eligible = _eligible_display_passages(passages, gate)

    assert [item["passage_id"] for item in eligible] == ["direct", "indirect"]


def test_indirect_passage_remains_when_it_is_the_only_relevant_fallback():
    passages = [
        {"passage_id": "indirect", "excerpt": "Another study reports the idea."},
    ]
    gate = {
        "status": "complete",
        "assessments": [
            {
                "passage_id": "indirect",
                "relevance": "partially_relevant",
                "evidence_role": "representation_of_other_work",
            },
        ],
    }

    eligible = _eligible_display_passages(passages, gate)

    assert [item["passage_id"] for item in eligible] == ["indirect"]


def test_partial_passages_remain_available_for_bounded_additional_context():
    passages = [
        {"passage_id": "first", "excerpt": "One partial connection."},
        {"passage_id": "second", "excerpt": "Another partial connection."},
    ]
    gate = {
        "status": "complete",
        "assessments": [
            {
                "passage_id": passage["passage_id"],
                "relevance": "partially_relevant",
                "evidence_role": "source_own_claim_or_finding",
            }
            for passage in passages
        ],
    }

    eligible = _eligible_display_passages(passages, gate)

    assert [item["passage_id"] for item in eligible] == ["first", "second"]


def test_display_priority_keeps_exact_check_evidence_ahead_of_shadow_relevance():
    passages = [
        {
            "passage_id": "semantic",
            "excerpt": "A semantically related source finding.",
            "retrieval_score": 0.99,
        },
        {
            "passage_id": "exact-quote",
            "excerpt": "The exact words quoted by the student.",
            "retrieval_score": 0.40,
        },
    ]
    gate = {
        "assessments": [
            {
                "passage_id": "semantic",
                "relevance": "relevant",
                "confidence": "high",
                "evidence_role": "source_own_claim_or_finding",
            },
            {
                "passage_id": "exact-quote",
                "relevance": "not_relevant",
                "confidence": "medium",
                "evidence_role": "unclear",
            },
        ]
    }

    ordered = _prioritize_display_passages(
        passages,
        "The student's quotation.",
        gate,
        preferred_passage_ids=["exact-quote"],
    )

    assert [item["passage_id"] for item in ordered] == [
        "exact-quote",
        "semantic",
    ]


def test_complete_relevance_gate_does_not_fill_report_with_rejected_passages():
    passages = [
        {"passage_id": "wrong-1", "excerpt": "Unrelated methods material."},
        {"passage_id": "wrong-2", "excerpt": "Unrelated literature review."},
    ]
    gate = {
        "status": "complete",
        "assessments": [
            {"passage_id": "wrong-1", "relevance": "not_relevant"},
            {"passage_id": "wrong-2", "relevance": "not_relevant"},
        ],
    }

    assert _eligible_display_passages(passages, gate) == []


def test_complete_relevance_gate_does_not_display_uncertain_candidate():
    passages = [
        {"passage_id": "uncertain", "excerpt": "Topically adjacent material."},
    ]
    gate = {
        "status": "complete",
        "assessments": [
            {"passage_id": "uncertain", "relevance": "uncertain"},
        ],
    }

    assert _eligible_display_passages(passages, gate) == []


def test_complete_relevance_gate_keeps_direct_check_evidence_when_rejected():
    passages = [
        {"passage_id": "exact-quote", "excerpt": "Exact quoted words."},
        {"passage_id": "wrong", "excerpt": "Unrelated passage."},
    ]
    gate = {
        "status": "complete",
        "assessments": [
            {"passage_id": "exact-quote", "relevance": "not_relevant"},
            {"passage_id": "wrong", "relevance": "not_relevant"},
        ],
    }

    assert _eligible_display_passages(
        passages,
        gate,
        preferred_passage_ids=["exact-quote"],
    ) == [passages[0]]


def test_incomplete_relevance_gate_retains_unassessed_retrieval_for_inspection():
    passages = [
        {"passage_id": "exact-quote", "excerpt": "Exact quoted words."},
        {"passage_id": "lexical", "excerpt": "Unreviewed lexical candidate."},
    ]

    assert _eligible_display_passages(passages, {"status": "not_assessed"}) == passages
    assert _eligible_display_passages(
        passages,
        {"status": "not_assessed"},
        preferred_passage_ids=["exact-quote"],
    ) == passages


def test_compound_citation_records_an_unavailable_member_without_borrowing(
    monkeypatch,
):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_compound_paper_bytes(),
            filename="partial-compound-paper.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
        )
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        assert retrieve_paper_sources(
            session,
            storage,
            job_id,
            resolver=PartialCompoundResolver(),
        ) == {"transient_authorized": 1, "unavailable": 1}

    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)

    assert summary["reports_persisted"] == 1
    assert summary["source_failures"] == [
        {
            "reference_id": summary["citation_groups"][0]["members"][1][
                "reference_id"
            ],
            "reason_code": "full_text_unavailable",
        }
    ]
    group = summary["citation_groups"][0]
    assert group["expected_member_count"] == 2
    assert group["persisted_member_count"] == 1
    assert group["coverage_status"] == "partial"
    assert [member["status"] for member in group["members"]] == [
        "evidence_package_persisted",
        "source_unavailable",
    ]
    assert group["members"][1]["reason_code"] == "full_text_unavailable"

    with factory() as session:
        reports = session.scalars(select(VerificationReportRecord)).all()
        assert len(reports) == 1
        package = reports[0].report_payload["authoritative_evidence_package"]
        assert package["source_binding"]["reference_id"] == group["members"][0][
            "reference_id"
        ]
    _assert_only_report_marking_copy(storage)


def test_provider_recovery_refreshes_only_affected_member_as_immutable_successor(
    monkeypatch,
):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_compound_paper_bytes(),
            filename="provider-recovery-paper.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
        )
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        retrieve_paper_sources(
            session, storage, job_id, resolver=PartialCompoundResolver()
        )
        source_results = json.loads(json.dumps(job.source_results))
        failed = next(item for item in source_results if item["status"] == "unavailable")
        failed_reference_id = failed["reference_id"]
        # The refresh recomputes blocking providers from the stored record, so
        # the record must genuinely be held incomplete by SearXNG. Under the
        # legacy policy SearXNG is a required web route, which it is not under
        # api-first.
        from tests.unit.blocked_discovery import blocked_record
        failed["reference_discovery"] = blocked_record(
            failed_reference_id, web_outcomes={"searxng": "cooldown_skipped"})
        failed["retryable_provider_dependencies"] = ["searxng"]
        job.source_results = source_results
        session.commit()

    first_summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert first_summary["reports_persisted"] == 1
    with factory() as session:
        first_report_id = finalize_paper_job(session, storage, job_id)["report_id"]
        assert prepare_provider_recovery_refresh(
            session, job_id, provider="searxng"
        ) == [failed_reference_id]
        resolver = CompoundResolver()
        retrieve_paper_sources(session, storage, job_id, resolver=resolver)
        assert resolver.calls == [failed_reference_id]

    refreshed = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert refreshed["reports_persisted"] == 2
    assert refreshed["source_failures"] == []
    assert refreshed["citation_groups"][0]["coverage_status"] == "complete"
    with factory() as session:
        second_report_id = finalize_paper_job(session, storage, job_id)["report_id"]
        assert second_report_id != first_report_id
        first_report = session.get(Report, uuid.UUID(first_report_id))
        second_report = session.get(Report, uuid.UUID(second_report_id))
        assert first_report.report_version == 1
        assert second_report.report_version == 2
        assert second_report.previous_report_id == first_report.id
        assert second_report.amendment_reason == "provider_recovery"
        assert len(session.scalars(select(Report)).all()) == 2
        assert len(session.scalars(select(VerificationReportRecord)).all()) == 2
        job = session.get(Job, job_id)
        assert "provider_refresh_reference_ids" not in job.upload_evidence
    _assert_only_report_marking_copy(storage)


def test_uploaded_source_reanalysis_is_targeted_idempotent_and_immutable(monkeypatch):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_compound_paper_bytes(with_link=True),
            filename="uploaded-source-refresh.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
        )
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        retrieve_paper_sources(
            session,
            storage,
            job_id,
            resolver=PartialCompoundResolver(),
        )
    verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        first_report_id = finalize_paper_job(session, storage, job_id)["report_id"]
        job = session.get(Job, job_id)
        extraction = job.extraction_payload
        jones = next(
            item for item in extraction["references"] if item["author"].startswith("Jones")
        )
        source_pdf = fitz.open()
        page = source_pdf.new_page()
        page.insert_text(
            (72, 72),
            "The Jones source independently establishes that the shared result is stable through repeated trials.",
        )
        source_bytes = source_pdf.tobytes()
        source_pdf.close()
        admitted = admit_representation(
            session,
            storage,
            AdmissionRequest(
                work=WorkIdentity(
                    title=jones["title"],
                    author=jones["author"],
                    year=jones["year"],
                    work_type="journal_article",
                ),
                representation=SourceRepresentation(
                    kind=RepresentationKind.PDF,
                    media_type="application/pdf",
                    content=source_bytes,
                ),
                provenance="instructor_upload",
                license_class="commercial_user_upload",
                scope_type=job.scope_type,
                scope_id=job.scope_id,
                identity_verdict="verified",
                identity_confidence=1.0,
                completeness_verdict="complete",
                cleanliness_verdict="clean",
                text_quality="digital",
            ),
        )
        commit_source_admissions(session)
        from app.services.submitted_links import initial_observations, LinkRequest
        from app.services.schemas import ParsedReference
        rows = initial_observations(ParsedReference.model_validate(jones))
        assert len(rows) == 1
        now = datetime.now(timezone.utc)
        rows[0].state = 'observed'
        rows[0].requests = [LinkRequest(started_at=now,completed_at=now,request_sha256=rows[0].request_sha256,http_status=404,outcome='not_found')]
        observed = [row.model_dump(mode='json') for row in rows]
        job.source_results = [{**item, 'submitted_link_observations': observed}
            if item.get('reference_id') == jones['reference_id'] else item for item in job.source_results]
        session.commit()
        prepared = prepare_uploaded_source_refresh(
            session,
            storage,
            report_id=first_report_id,
            reference_id=jones["reference_id"],
            representation_id=admitted.id,
        )
        assert prepared["scheduled"] is True
        replacement = next(item for item in job.source_results if item['reference_id'] == jones['reference_id'])
        assert replacement['submitted_link_observations'] == observed
        assert session.get(Job, uuid.UUID(prepared["job_id"])).upload_evidence["workflow_dispatch_v1"]["attempt_id"] == prepared["attempt_id"]
        duplicate = prepare_uploaded_source_refresh(
            session,
            storage,
            report_id=first_report_id,
            reference_id=jones["reference_id"],
            representation_id=admitted.id,
        )
        assert duplicate["status"] == "already_scheduled"

    refreshed = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert refreshed["reports_persisted"] == 2
    assert refreshed["source_failures"] == []
    with factory() as session:
        second_report_id = finalize_paper_job(session, storage, job_id)["report_id"]
        first = session.get(Report, uuid.UUID(first_report_id))
        second = session.get(Report, uuid.UUID(second_report_id))
        assert first.report_version == 1
        assert second.report_version == 2
        assert second.previous_report_id == first.id
        assert second.amendment_reason == "user_source_upload"
        amendment = second.report_json["amendment"]
        assert amendment["base_report_id"] == first_report_id
        assert amendment["reference_ids"] == [jones["reference_id"]]
        assert len(amendment["changed_dependencies"]) == 1
        assert amendment["changed_dependencies"][0]["old_status"] == "source_unavailable"
        assert amendment["changed_dependencies"][0]["new_status"] == "evidence_package_persisted"
        assert len(session.scalars(select(Report)).all()) == 2
        already_current = prepare_uploaded_source_refresh(
            session,
            storage,
            report_id=first_report_id,
            reference_id=jones["reference_id"],
            representation_id=admitted.id,
        )
        assert already_current["status"] == "already_current"
        assert already_current["report_id"] == second_report_id


def test_targeted_source_refresh_failure_restores_previous_checkpoint(monkeypatch):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_compound_paper_bytes(),
            filename="rollback-source-refresh.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
        )
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        retrieve_paper_sources(
            session,
            storage,
            job_id,
            resolver=PartialCompoundResolver(),
        )
    verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    with factory() as session:
        report_id = finalize_paper_job(session, storage, job_id)["report_id"]
        job = session.get(Job, job_id)
        original_results = json.loads(json.dumps(job.source_results))
        original_summary = json.loads(json.dumps(job.verification_summary))
        jones = next(
            item
            for item in job.extraction_payload["references"]
            if item["author"].startswith("Jones")
        )
        admitted = admit_representation(
            session,
            storage,
            AdmissionRequest(
                work=WorkIdentity(
                    title=jones["title"],
                    author=jones["author"],
                    year=jones["year"],
                    work_type="journal_article",
                ),
                representation=SourceRepresentation(
                    kind=RepresentationKind.PLAIN_TEXT,
                    media_type="text/plain",
                    content=b"A complete alternate source representation.",
                ),
                provenance="instructor_upload",
                license_class="commercial_user_upload",
                scope_type=job.scope_type,
                scope_id=job.scope_id,
                identity_verdict="verified",
                identity_confidence=1.0,
                completeness_verdict="complete",
                cleanliness_verdict="clean",
                text_quality="digital",
            ),
        )
        commit_source_admissions(session)
        prepare_uploaded_source_refresh(
            session,
            storage,
            report_id=report_id,
            reference_id=jones["reference_id"],
            representation_id=admitted.id,
        )
        assert rollback_targeted_source_refresh(
            session, job_id, error_code="simulated_downstream_failure"
        )
        restored = session.get(Job, job_id)
        assert restored.status == "completed"
        assert restored.stage == "completed"
        assert restored.source_results == original_results
        assert restored.verification_summary == original_summary
        assert len(session.scalars(select(Report)).all()) == 1
        retried = prepare_uploaded_source_refresh(
            session,
            storage,
            report_id=report_id,
            reference_id=jones["reference_id"],
            representation_id=admitted.id,
        )
        assert retried["scheduled"] is True


def test_compound_citation_keeps_member_evidence_packages_source_separate(
    monkeypatch,
):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    resolver = CompoundResolver()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_compound_paper_bytes(),
            filename="compound-paper.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
        )
        job_id = job.id
        counts = extract_paper_job(session, storage, job_id, llm_enabled=False)
        assert counts["references"] == 2
        assert counts["citation_units"] == 1
        extraction = job.extraction_payload
        claim = extraction["citation_claims"][0]
        assert len(claim["reference_ids"]) == 2
        assert [member["text"] for member in claim["citation_markers"]] == [
            "Smith (2020)",
            "Jones (2021)",
        ]
        assert retrieve_paper_sources(
            session, storage, job_id, resolver=resolver
        ) == {"transient_authorized": 2}
        assert len(set(resolver.calls)) == 2

    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    assert summary["reports_persisted"] == 2
    assert len(summary["citation_groups"]) == 1
    group = summary["citation_groups"][0]
    assert group["expected_member_count"] == 2
    assert group["persisted_member_count"] == 2
    assert group["coverage_status"] == "complete"
    assert {member["reference_id"] for member in group["members"]} == set(
        claim["reference_ids"]
    )
    assert all(
        member["status"] == "evidence_package_persisted"
        for member in group["members"]
    )

    with factory() as session:
        reports = session.scalars(
            select(VerificationReportRecord).order_by(
                VerificationReportRecord.created_at
            )
        ).all()
        runs = session.scalars(select(VerificationRunRecord)).all()
        assert len(reports) == 2
        assert len(runs) == 2
        assert all(run.status == "cleaned" for run in runs)

        packages = [
            report.report_payload["authoritative_evidence_package"]
            for report in reports
        ]
        assert len({package["package_id"] for package in packages}) == 2
        assert len(
            {package["source_binding"]["reference_id"] for package in packages}
        ) == 2
        assert {package["source_binding"]["marker_text"] for package in packages} == {
            "Smith (2020)",
            "Jones (2021)",
        }
        assert len(
            {package["source_identity"]["content_sha256"] for package in packages}
        ) == 2
        assert len({package["claim_id"] for package in packages}) == 1
        for package in packages:
            author = package["source_binding"]["cited_author_label"].split(",", 1)[0]
            excerpts = " ".join(
                passage["excerpt"] for passage in package["passages"]
            )
            assert author in excerpts
            other = "Jones" if author == "Smith" else "Smith"
            assert other not in excerpts
    _assert_only_report_marking_copy(storage)


@pytest.mark.parametrize("through_tasks", [False, True])
def test_store_only_job_stops_after_extraction_and_cleans(monkeypatch, through_tasks):
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(
            session,
            storage,
            content=_paper_bytes(),
            filename="paper.docx",
            media_type=DOCX_MEDIA_TYPE,
            scope_id="personal-default",
            store_only=True,
        )
        job_id = job.id
        if through_tasks:
            from app.tasks import check_paper as tasks
            from app.config import settings
            from app.services.paper_dispatch import attempt_id_for
            monkeypatch.setattr(settings, "PAPER_LLM_PROCESSING_ENABLED", False)
            monkeypatch.setattr(tasks, "SessionLocal", factory)
            monkeypatch.setattr(tasks, "get_storage_backend", lambda: storage)
            attempt = attempt_id_for(job)
            tasks.extract_paper_task.run(str(job_id), attempt)
            finalized = tasks.finalize_paper_job_task.run(str(job_id), attempt)
            session.expire_all()
        else:
            extract_paper_job(session, storage, job_id, llm_enabled=False)
            finalized = finalize_paper_job(session, storage, job_id)
        aggregate = session.scalar(select(Report))
        assert finalized["input_cleaned"] is True
        assert aggregate.report_json["store_only"] is True
        assert aggregate.report_json["reports_persisted"] == 0
        assert session.scalar(select(VerificationRunRecord)) is None
        assert session.scalar(select(VerificationReportRecord)) is None
    _assert_only_report_marking_copy(storage)


def test_patchwriting_runs_inside_the_paper_check_and_is_not_rendered(monkeypatch):
    # Owner decision 2026-09-29: patchwriting-v2 runs while each source's text is
    # authorized, including a source held only for the run; nothing is rendered.
    from app.services import patchwriting as pw
    monkeypatch.setattr(paper_upload, 'scan_with_clamd',
        lambda _: (SafetyVerdict.CLEAN, 'stream: OK'))
    from app.config import settings
    monkeypatch.setattr(settings, 'PATCHWRITING_AT_CHECK_ENABLED', True)
    engine = create_engine('sqlite+pysqlite:///:memory:')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    storage = MemoryStorage()
    with factory() as session:
        job = create_paper_job(session, storage, content=_paper_bytes(),
            filename='paper.docx', media_type=DOCX_MEDIA_TYPE, scope_id='personal-default')
        job_id = job.id
        extract_paper_job(session, storage, job_id, llm_enabled=False)
        retrieve_paper_sources(session, storage, job_id, resolver=ProvisionalResolver())
    summary = verify_paper_sources(factory, storage, job_id, llm_enabled=False)
    block = summary['patchwriting']
    assert block['policy_version'] == pw.POLICY_VERSION and block['decision_applied'] is False
    assert block['sources'], 'the run-only source was compared'
    assert block['body']['status'] == 'ready', block['body']
    assert all(v.get('status') != 'not_assessed' for v in block['sources'].values()), block['sources']
    for entry in block['sources'].values():
        assert entry['policy_version'] == pw.POLICY_VERSION
    with factory() as session:
        finalize_paper_job(session, storage, job_id)
        view = session.scalar(select(Report)).report_json['evidence_report']
    html = render_evidence_report_html(view, csp_nonce='test-patchwriting-nonce')
    assert 'close_paraphrase' not in html and 'unquoted_verbatim' not in html

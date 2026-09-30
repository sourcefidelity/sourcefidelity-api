import io
import pytest
from docx import Document
from app.services.schemas import ParsedReference
from app.services.body_title_formatting import assess_body_title_italics


@pytest.mark.parametrize('mode,expected',[('plain','difference'),('italic','matches_rule'),
    ('ambiguous',None),('no_year',None),('wrong_year',None),('whole_italic','not_assessed'),
    ('reference_only',None),('review',None),('mla',None)])
def test_exact_title_year_style_with_conservative_boundaries(mode,expected):
    d=Document();p=d.add_paragraph();p.add_run('In ')
    title=p.add_run('Example Film');title.italic=mode in {'italic','whole_italic'}
    p.add_run(' (1991) the staging changes.' if mode!='no_year' else ' the staging changes.')
    if mode=='wrong_year':p.runs[-1].text=' (1992) the staging changes.'
    if mode=='whole_italic':
        for r in p.runs:r.italic=True
    body=p.text
    if mode=='reference_only':body='Different body.'
    ref=ParsedReference(reference_id='film',title='Example Film',year='1991',raw_ref='Author. Example Film [Film].',
        source_kind='traditional_media',source_kind_confidence='high',needs_review=mode=='review')
    refs=[ref,ref.model_copy(update={'reference_id':'second'})] if mode=='ambiguous' else [ref]
    stream=io.BytesIO();d.save(stream)
    result=assess_body_title_italics(content=stream.getvalue(),body=body,references=refs,citation_format='mla' if mode=='mla' else 'apa')
    assert [o['status'] for o in result['observations']]==([] if expected is None else [expected])
    assert not result['automatic_findings_enabled']

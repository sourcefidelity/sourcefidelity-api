import hashlib
import io
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from app.services.reference_layout import (
    ReferenceEntryLayout, _LayoutLine, _bind_style_evidence,
    _docx_run_italic, extract_reference_layout_from_bytes, reference_span_style_observed,
)
from app.services.schemas import ParsedReference


def test_plain_text_is_observed_but_historical_or_changed_text_is_not():
    raw='Smith, J. (2020). A plain title. Publisher.'
    doc=Document(); doc.add_paragraph('References');doc.add_paragraph(raw)
    output=io.BytesIO();doc.save(output)
    entry=extract_reference_layout_from_bytes(output.getvalue(),'case.docx',
        references=[ParsedReference(reference_id='r',raw_ref=raw)],citation_format='apa').entries[0]
    assert entry.text_style_spans == []
    assert reference_span_style_observed(entry,raw,0,len(raw))
    assert not reference_span_style_observed(entry,raw+'changed',0,len(raw))
    historical=entry.model_dump();historical.pop('style_binding_version');historical.pop('style_observation_ranges')
    assert not reference_span_style_observed(ReferenceEntryLayout.model_validate(historical),raw,0,len(raw))
    assert ReferenceEntryLayout.model_validate_json(entry.model_dump_json()) == entry


def test_unmatched_characters_are_not_silently_unitalicized():
    raw='An altered title'
    line=_LayoutLine(text='An original title',location_index=0,x0=0,continuation_x0=None,
        italic_characters=0,bold_characters=0,styled_characters=17)
    styles,ranges=_bind_style_evidence(raw,[line])
    entry=ReferenceEntryLayout(reference_id='r',reference_text_sha256=hashlib.sha256(raw.encode()).hexdigest(),
        mapping_status='matched',match_confidence=.8,line_count=1,
        style_binding_version='reference-style-binding-v1',style_observation_ranges=ranges)
    assert not styles
    assert not reference_span_style_observed(entry,raw,0,len(raw))
    assert reference_span_style_observed(entry,raw,raw.index('title'),len(raw))


@pytest.mark.parametrize('direct,expected', [(None,True),(False,False),(True,True)])
def test_paragraph_style_inheritance_and_direct_override(direct,expected):
    doc=Document()
    base=doc.styles.add_style('BaseItalic',WD_STYLE_TYPE.PARAGRAPH);base.font.italic=True
    child=doc.styles.add_style('InheritedItalic',WD_STYLE_TYPE.PARAGRAPH);child.base_style=base
    run=doc.add_paragraph(style=child).add_run('A title');run.italic=direct
    assert _docx_run_italic(run) is expected


def test_character_style_toggle_and_base_style_cycle():
    doc=Document()
    base=doc.styles.add_style('BaseItalic',WD_STYLE_TYPE.CHARACTER);base.font.italic=True
    child=doc.styles.add_style('CancelItalic',WD_STYLE_TYPE.CHARACTER);child.base_style=base;child.font.italic=True
    run=doc.add_paragraph().add_run('A title');run.style=child
    assert _docx_run_italic(run) is False
    base.base_style=child
    assert _docx_run_italic(run) is None


def test_document_default_italic_is_observed():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    defaults=doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    defaults.append(OxmlElement('w:i'))
    run=doc.add_paragraph().add_run('A title')
    assert _docx_run_italic(run) is True


def test_soft_break_references_preserve_styles_without_inventing_indentation():
    refs=['Smith, J. (2020). First book. Publisher.',
          'Jones, A. (2021). Second book. Other publisher.']
    doc=Document();doc.add_paragraph('References');p=doc.add_paragraph()
    p.add_run(refs[0]).italic=True
    p.add_run('\n\n')
    p.add_run(refs[1])
    output=io.BytesIO();doc.save(output)
    artifact=extract_reference_layout_from_bytes(output.getvalue(),'case.docx',
        references=[ParsedReference(reference_id=str(i),raw_ref=r) for i,r in enumerate(refs)],
        citation_format='apa')
    assert artifact.matched_reference_count == 2
    for raw,entry in zip(refs,artifact.entries):
        assert reference_span_style_observed(entry,raw,0,len(raw))
        assert entry.observed_hanging_indent_points is None
        assert entry.location_indexes == [1]
    assert artifact.entries[0].text_style_spans
    assert not artifact.entries[1].text_style_spans

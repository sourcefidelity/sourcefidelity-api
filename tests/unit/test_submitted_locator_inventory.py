import hashlib
import io
import json

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.services.schemas import ParsedReference
from app.services.reference_layout import extract_reference_layout_from_bytes
from app.services.submitted_locator_inventory import inventory_submitted_locators, SubmittedLocatorInventory


def paper(tail='', *, hidden_link=False, style='apa'):
    raw='Smith, J. (2020). A bounded title. Publisher. '+tail
    with fitz.open() as doc:
        page=doc.new_page()
        page.insert_text((72,72),'References' if style=='apa' else 'Works Cited')
        page.insert_text((72,110),raw,fontsize=9)
        if hidden_link:
            page.insert_link({'kind':fitz.LINK_URI,'from':fitz.Rect(72,100,140,114),
                              'uri':'https://source.example/item?token=PRIVATE_SENTINEL'})
        content=doc.tobytes()
    refs=[ParsedReference(reference_id='r1',raw_ref=raw)]
    layout=extract_reference_layout_from_bytes(content,'paper.pdf',references=refs,citation_format=style)
    return content,refs,layout


@pytest.mark.parametrize('tail',['10.1234/example','https://doi.org/10.1234/example',
                                  'https://source.example/a','https://','doi: broken',
                                  'www.example.org/work','example.org/work','https: //source.example/a'])
def test_supplied_including_malformed_is_never_absent(tail):
    content,refs,layout=paper(tail)
    result=inventory_submitted_locators(content,references=refs,layout=layout)
    assert result.counts=={'supplied':1,'not_observed':0,'unknown':0}
    assert result.entries[0].locator_validity=='not_assessed'


def test_native_link_counts_without_exposing_target():
    content,refs,layout=paper(hidden_link=True)
    result=inventory_submitted_locators(content,references=refs,layout=layout)
    assert result.entries[0].status=='supplied'
    assert 'native_hyperlink' in result.entries[0].evidence_channels
    assert 'PRIVATE_SENTINEL' not in result.model_dump_json()
    assert 'source.example' not in result.model_dump_json()


@pytest.mark.parametrize('style',['apa','mla'])
def test_absence_is_neutral_and_enrichment_does_not_change_it(style):
    content,refs,layout=paper(style=style)
    result=inventory_submitted_locators(content,references=refs,layout=layout)
    changed=[refs[0].model_copy(update={'doi':'10.9999/enriched','url':'https://enriched.example'})]
    later=inventory_submitted_locators(content,references=changed,layout=layout)
    assert result==later
    assert result.counts['not_observed']==1
    assert result.reference_list_coverage=='unknown'
    assert result.entries[0].rectangles and not result.contributes_to_issue_counts
    assert result.entries[0].style_requirement=='not_assessed'
    assert SubmittedLocatorInventory.model_validate_json(result.model_dump_json())==result


def test_uncertain_extraction_not_counted_as_absence():
    content,refs,layout=paper()
    refs[0].needs_review=True
    result=inventory_submitted_locators(content,references=refs,layout=layout)
    assert result.counts['unknown']==1 and result.assessed_entries==0


def test_snapshot_and_id_changes_fail():
    content,refs,layout=paper()
    with pytest.raises(ValueError):inventory_submitted_locators(content+b'changed',references=refs,layout=layout)
    with pytest.raises(ValueError):inventory_submitted_locators(content,references=refs*2,layout=layout)
    refs[0].raw_ref+='Changed'
    with pytest.raises(ValueError):inventory_submitted_locators(content,references=refs,layout=layout)


def test_count_tampering_rejected():
    content,refs,layout=paper()
    value=inventory_submitted_locators(content,references=refs,layout=layout).model_dump()
    value['counts']['not_observed']=99
    with pytest.raises(ValueError):SubmittedLocatorInventory.model_validate(value)


def test_visible_wrapped_url_is_not_lost():
    with fitz.open() as doc:
        page=doc.new_page();page.insert_text((72,72),'References')
        page.insert_text((72,110),'Smith, J. (2020). A bounded title. Publisher. https://')
        page.insert_text((72,130),'example.org/long/wrapped-path')
        content=doc.tobytes()
    raw='Smith, J. (2020). A bounded title. Publisher. https:// example.org/long/wrapped-path'
    refs=[ParsedReference(reference_id='r1',raw_ref=raw)]
    layout=extract_reference_layout_from_bytes(content,'paper.pdf',references=refs,citation_format='apa')
    result=inventory_submitted_locators(content,references=refs,layout=layout)
    assert result.entries[0].status=='supplied'


def test_initials_and_prose_do_not_become_bare_domains():
    from app.services.submitted_locator_inventory import _has_locator
    assert not _has_locator('J.Smith. Book.Title. Publisher.Name. user@example.org')


def test_page_transition_does_not_claim_absence_or_attribute_neighbor_link():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72,72), 'References')
        page.insert_text((72,110), 'Smith, J. (2020). A bounded title. Publisher.')
        page = doc.new_page()
        page.insert_text((72,72), 'Jones, J. (2021). Another title. https://example.org')
        content = doc.tobytes()
    refs = [ParsedReference(reference_id='a', raw_ref='Smith, J. (2020). A bounded title. Publisher.'),
            ParsedReference(reference_id='b', raw_ref='Jones, J. (2021). Another title. https://example.org')]
    layout = extract_reference_layout_from_bytes(content, 'paper.pdf', references=refs, citation_format='apa')
    result = inventory_submitted_locators(content, references=refs, layout=layout)
    assert result.counts['unknown'] == 2
    assert all(e.limitation == 'cross_page_entry_boundary_unverified' for e in result.entries)


def test_merged_numbered_entries_abstain():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72,72), 'Works Cited')
        page.insert_text((72,110), '1. Smith, J. A bounded title. Publisher, 2020.')
        page.insert_text((72,130), '2. Jones, J. Another title. Publisher, 2021.')
        content = doc.tobytes()
    refs = [ParsedReference(reference_id='a', raw_ref='1. Smith, J. A bounded title. Publisher, 2020. 2. Jones, J. Another title. Publisher, 2021.')]
    layout = extract_reference_layout_from_bytes(content, 'paper.pdf', references=refs, citation_format='mla')
    result = inventory_submitted_locators(content, references=refs, layout=layout)
    assert result.entries[0].status == 'unknown'


@pytest.mark.parametrize('kind,expected', [('plain','not_observed'), ('link','supplied'),
    ('broken','unknown'), ('field','unknown'), ('revision','unknown'), ('internal','not_observed')])
def test_docx_native_relationships(kind, expected):
    document = Document()
    document.add_paragraph('References')
    raw = 'Smith, J. (2020). A bounded title. Publisher.'
    p = document.add_paragraph(raw)
    if kind in ('link', 'broken', 'internal'):
        link = OxmlElement('w:hyperlink')
        if kind == 'internal':
            link.set(qn('w:anchor'), 'bookmark')
        else:
            rid = p.part.relate_to('https://example.org/PRIVATE_SENTINEL', RT.HYPERLINK, is_external=True)
            link.set(qn('r:id'), rid if kind == 'link' else 'rIdMissing')
        p._p.append(link)
    elif kind == 'field':
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'HYPERLINK "https://example.org"')
        p._p.append(field)
    elif kind == 'revision':
        p._p.append(OxmlElement('w:del'))
    stream = io.BytesIO(); document.save(stream); content = stream.getvalue()
    refs = [ParsedReference(reference_id='a', raw_ref=raw)]
    layout = extract_reference_layout_from_bytes(content, 'paper.docx', references=refs, citation_format='apa')
    result = inventory_submitted_locators(content, references=refs, layout=layout)
    assert result.entries[0].status == expected
    assert result.entries[0].paragraph_indexes == (1,)
    assert 'PRIVATE_SENTINEL' not in result.model_dump_json()
    enriched = [refs[0].model_copy(update={'url':'https://enriched.example'})]
    assert result == inventory_submitted_locators(content, references=enriched, layout=layout)


def test_docx_visible_hyperlink_runs_participate_in_layout():
    document = Document()
    document.add_paragraph('References')
    base = 'Smith, J. (2020). A bounded title. Publisher. '
    url = 'https://example.org/a-long-visible-locator'
    p = document.add_paragraph(base)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = url
    run.append(text); link.append(run); p._p.append(link)
    stream = io.BytesIO(); document.save(stream); content = stream.getvalue()
    refs = [ParsedReference(reference_id='a', raw_ref=base + url)]
    layout = extract_reference_layout_from_bytes(content, 'paper.docx', references=refs, citation_format='apa')
    assert layout.entries[0].match_confidence == 1
    entry = inventory_submitted_locators(content, references=refs, layout=layout).entries[0]
    assert entry.status == 'supplied'
    assert set(entry.evidence_channels) == {'raw_text','visible_text','native_hyperlink'}

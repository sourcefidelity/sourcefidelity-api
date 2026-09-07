"""Personal-profile paper submission and bounded status API."""

import io
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
import pytest
import uuid

from app.database import get_db
from app.config import settings
from app.main import app
from app.models import Base
from app.models.job import Job
from app.services.paper_dispatch import DISPATCH_KEY
from app.routers import check
from app.services import paper_upload
from app.services.file_safety import SafetyVerdict
from app.services.storage.backend import StorageBackend, get_storage_backend


class MemoryStorage(StorageBackend):
    def __init__(self):
        self.objects = {}

    def upload(self, file_bytes, key):
        self.objects[key] = file_bytes
        return key

    def download(self, key):
        if key not in self.objects:
            raise FileNotFoundError(key)
        return self.objects[key]

    def delete(self, key):
        self.objects.pop(key, None)
        return True

    def exists(self, key):
        return key in self.objects

    def list_keys(self, prefix):
        return [key for key in self.objects if key.startswith(prefix)]


def _docx_bytes():
    document = Document()
    document.add_paragraph("A citation (Smith, 2020).")
    document.add_paragraph("References")
    document.add_paragraph("Smith, J. (2020). A useful title. Example Press.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("queue_available", [True, False])
def test_submit_persists_recoverable_job_and_status_is_shadow_only(monkeypatch, queue_available):
    monkeypatch.setattr(settings, "REPORT_AUTH_MODE", "personal_local")
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    storage = MemoryStorage()

    def override_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_storage_backend] = lambda: storage
    monkeypatch.setattr(
        paper_upload,
        "scan_with_clamd",
        lambda _content: (SafetyVerdict.CLEAN, "stream: OK"),
    )
    def publish(job_id, attempt_id):
        with Session(engine) as session:
            assert session.get(Job, uuid.UUID(job_id)).upload_evidence[DISPATCH_KEY]["attempt_id"] == attempt_id
        if not queue_available:
            raise ConnectionError("queue unavailable")
        return SimpleNamespace(id="queued-task-1")

    monkeypatch.setattr(
        check.check_paper_task,
        "delay",
        publish,
    )
    try:
        client = TestClient(app)
        response = client.post(
            "/check/",
            files={
                "file": (
                    "paper.docx",
                    _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 202
        submitted = response.json()
        assert submitted["status"] == "pending"
        assert submitted["stage"] == "uploaded"
        assert submitted["publication_pending"] is (not queue_available)
        assert len(storage.objects) == 1

        status_response = client.get(f"/status/{submitted['job_id']}")
        assert status_response.status_code == 200
        status_payload = status_response.json()
        assert status_payload["reports_persisted"] == 0
        assert status_payload["decision_applied"] is False
        assert status_payload["error_code"] is None
    finally:
        app.dependency_overrides.clear()
